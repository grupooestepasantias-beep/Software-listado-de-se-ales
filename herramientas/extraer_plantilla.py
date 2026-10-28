"""Extrae la plantilla de formato de un listado de señales de ejemplo (solo lectura).

Uso:
    python herramientas/extraer_plantilla.py <ejemplo.xlsx> <salida.json> [hoja_caratula] [hoja_listado]

Guarda anchos, alturas, merges, estilos, filtros, área de impresión, títulos y
formato condicional. No guarda datos de señales: solo formato y rótulos fijos.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

FILAS_ENCABEZADO = 5   # filas 1-5 del listado
FILAS_MUESTRA = (6, 7, 8)  # filas de datos usadas como muestra de estilo


def _color(c):
    if c is None:
        return None
    if c.type == "rgb":
        return c.rgb if isinstance(c.rgb, str) else None
    if c.type == "theme":
        return {"theme": c.theme, "tint": c.tint}
    if c.type == "indexed":
        return {"indexed": c.indexed}
    return None


def _lado(s):
    if s is None or s.style is None:
        return None
    return {"style": s.style, "color": _color(s.color)}


def estilo(cell) -> dict:
    f, fl, b, a = cell.font, cell.fill, cell.border, cell.alignment
    d = {
        "font": {"name": f.name, "sz": f.sz, "b": bool(f.b), "i": bool(f.i),
                 "u": f.u, "color": _color(f.color)},
        "fill": None,
        "border": {k: _lado(getattr(b, k)) for k in ("left", "right", "top", "bottom")},
        "align": {"h": a.horizontal, "v": a.vertical, "wrap": bool(a.wrap_text),
                  "rot": a.textRotation, "indent": a.indent},
        "fmt": cell.number_format,
    }
    if fl is not None and fl.fill_type:
        d["fill"] = {"type": fl.fill_type, "fg": _color(fl.fgColor)}
    return d


def _celda(cell) -> dict:
    d = estilo(cell)
    d["v"] = cell.value
    return d


def _vacia(cell) -> bool:
    e = estilo(cell)
    sin_borde = all(v is None for v in e["border"].values())
    return cell.value is None and e["fill"] is None and sin_borde


def extraer_listado(ws) -> dict:
    ult = ws.max_column
    columnas = {}
    for i in range(1, ult + 1):
        l = get_column_letter(i)
        dim = ws.column_dimensions.get(l)
        columnas[l] = {"ancho": dim.width if dim else None,
                       "oculta": bool(dim.hidden) if dim else False}
    alturas = {str(r): ws.row_dimensions[r].height
               for r in list(range(1, FILAS_ENCABEZADO + 1)) + list(FILAS_MUESTRA)
               if ws.row_dimensions[r].height is not None}
    encabezado = {str(r): {get_column_letter(c): _celda(ws.cell(r, c))
                           for c in range(1, ult + 1)}
                  for r in range(1, FILAS_ENCABEZADO + 1)}
    muestra = {str(r): {get_column_letter(c): estilo(ws.cell(r, c))
                        for c in range(1, ult + 1)}
               for r in FILAS_MUESTRA}
    merges_enc = [str(m) for m in ws.merged_cells.ranges if m.min_row <= FILAS_ENCABEZADO]
    cf = []
    for rng, reglas in ws.conditional_formatting._cf_rules.items():
        for r in reglas:
            cf.append({"rango": str(rng.sqref), "tipo": r.type,
                       "operador": r.operator, "formula": list(r.formula or []),
                       "dxf_fill": _color(r.dxf.fill.bgColor) if r.dxf and r.dxf.fill else None,
                       "dxf_font": _color(r.dxf.font.color) if r.dxf and r.dxf.font else None})
    return {
        "dimension": ws.dimensions,
        "columnas": columnas,
        "alturas_filas": alturas,
        "encabezado": encabezado,
        "merges_encabezado": merges_enc,
        "estilo_filas_datos": muestra,
        "autofiltro": ws.auto_filter.ref,
        "area_impresion": ws.print_area,
        "titulos_repetir_filas": ws.print_title_rows,
        "titulos_repetir_columnas": ws.print_title_cols,
        "congelar": ws.freeze_panes,
        "formato_condicional": cf,
        "pagina": {"orientacion": ws.page_setup.orientation,
                   "papel": ws.page_setup.paperSize,
                   "escala": ws.page_setup.scale,
                   "ajustar_ancho": ws.page_setup.fitToWidth,
                   "ajustar_alto": ws.page_setup.fitToHeight,
                   "fit_to_page": ws.sheet_properties.pageSetUpPr.fitToPage
                   if ws.sheet_properties.pageSetUpPr else None},
        "margenes": {k: getattr(ws.page_margins, k)
                     for k in ("left", "right", "top", "bottom", "header", "footer")},
    }


def extraer_caratula(ws) -> dict:
    columnas = {}
    for i in range(1, ws.max_column + 1):
        l = get_column_letter(i)
        dim = ws.column_dimensions.get(l)
        columnas[l] = dim.width if dim else None
    alturas = {str(r): ws.row_dimensions[r].height
               for r in range(1, ws.max_row + 1) if ws.row_dimensions[r].height is not None}
    celdas = {}
    for fila in ws.iter_rows():
        for c in fila:
            if not _vacia(c):
                celdas[c.coordinate] = _celda(c)
    return {
        "dimension": ws.dimensions,
        "columnas": columnas,
        "alturas_filas": alturas,
        "merges": sorted(str(m) for m in ws.merged_cells.ranges),
        "celdas": celdas,
        "area_impresion": ws.print_area,
        "pagina": {"orientacion": ws.page_setup.orientation,
                   "papel": ws.page_setup.paperSize,
                   "escala": ws.page_setup.scale},
        "margenes": {k: getattr(ws.page_margins, k)
                     for k in ("left", "right", "top", "bottom", "header", "footer")},
    }


# Dónde va cada dato dinámico de la carátula. Es formato (posición de celdas), por eso vive en la
# plantilla y no en el código del escritor. Medido sobre el ejemplo ETAP (EPEN).
CAMPOS_CARATULA = {
    "referencias": {"filas": [3, 16], "col_codigo": "B", "col_titulo": "G"},
    # la revisión más reciente queda arriba; las filas se llenan de abajo hacia arriba
    "revisiones": {"filas": [25, 28],
                   "columnas": {"rev": "A", "fecha": "C", "descripcion": "I",
                                "proyecto": "V", "revisor": "Y", "aprobador": "AB"}},
    "simples": {
        "cliente": {"celda": "C31"},
        "obra": {"celda": "C34"},
        "lugar": {"celda": "A38", "formato": "   Lugar: {}"},
        "departamento": {"celda": "H38", "formato": "   Depto.: {}"},
        "expediente": {"celda": "S38", "formato": "   Nro. Expediente: {}"},
        "titulo": {"celda": "A41"},
        "proyecto": {"celda": "C46"},
        "dibujo": {"celda": "C48"},
        "aprobo": {"celda": "C50"},
        "archivo_magnetico": {"celda": "J47"},
        "escala": {"celda": "R46"},
        "revision": {"celda": "X46"},
        "fecha": {"celda": "J50"},
    },
    # un carácter por celda, sin guiones (los separadores son celdas vacías / el "-" fijo de N44)
    "documento": {"celdas": ["D44", "E44", "F44", "G44", "I44", "J44", "K44", "L44",
                             "O44", "P44", "R44", "S44", "U44", "V44", "X44", "Y44",
                             "AA44", "AB44", "AC44", "AD44"]},
    "hoja": {"celda": "T49", "formato": "{n} de {total}"},
}


def extraer(ruta, hoja_caratula="Caratula", hoja_listado="EPEN") -> dict:
    wb = load_workbook(ruta)  # sin data_only: no hacen falta valores calculados
    return {
        "caratula_campos": CAMPOS_CARATULA,
        "fuente": Path(ruta).name,
        "hoja_caratula": hoja_caratula,
        "hoja_listado": hoja_listado,
        "caratula": extraer_caratula(wb[hoja_caratula]),
        "listado": extraer_listado(wb[hoja_listado]),
    }


def main(argv: list[str]) -> int:
    if len(argv) < 3:
        print(__doc__)
        return 1
    datos = extraer(argv[1], *argv[3:5])
    out = Path(argv[2])
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(datos, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print(f"Plantilla guardada en {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
