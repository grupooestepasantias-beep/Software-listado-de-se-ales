"""Extrae las abreviaturas de DESC. SCADA de un listado ya hecho (Excel) y las guarda como reglas del cliente.

Uso:  python herramientas/extraer_reglas.py <listado.xlsx> [id_perfil]

Escribe/actualiza plantillas/<id>/reglas_descripcion.json (equipo, frase, estado, texto). Donde el Excel
y los lineamientos difieren, gana lo que hace el Excel (es lo que usa la empresa); lo que el Excel no
tiene se conserva tal cual. Si una abreviatura aparece con valores distintos, se toma la más usada y
queda anotada en "_conflictos" para revisarla.
"""
from __future__ import annotations

import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import descripciones as dsc  # noqa: E402
import perfiles  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _sin_prefijo(hmi: str, scada: str):
    pref = hmi.split(" ", 1)[0]
    if not pref or not scada.startswith(pref):
        return None
    return scada[len(pref):].strip()


def extraer(ruta, perfil) -> dict:
    letra = {c["clave"]: c["letra"] for c in perfil.columnas}
    ws = load_workbook(ruta, read_only=True, data_only=True)[perfil.datos["hojas"]["listado"]]
    filas = []
    for r in ws.iter_rows(min_row=perfil.datos["hojas"]["primera_fila_datos"], values_only=True):
        g = lambda k: r[ord(letra[k]) - 65] if ord(letra[k]) - 65 < len(r) else None  # noqa: E731
        hmi, scada, obs = (str(g(k) or "").strip() for k in ("desc_hmi", "desc_scada", "observaciones"))
        if not (hmi and scada and obs) or str(g("tag_scada") or "").upper() == "N/A":
            continue
        suf = _sin_prefijo(hmi, scada)
        if suf:
            filas.append((obs, suf, (g("s0"), g("s1"), g("s2"), g("s3")), str(g("funcion") or "")))

    sin = json.loads((dsc.RAIZ_PLANTILLAS / perfil.id / "reglas_descripcion.json")
                     .read_text(encoding="utf-8")).get("sinonimos_equipo", {})
    por_equipo, con_equipo, sin_equipo = defaultdict(list), [], []
    for obs, suf, est, fn in filas:
        eq, letras, frase = dsc.analizar(obs, sin)
        (con_equipo if eq else sin_equipo).append((eq, letras, frase, suf, est, obs))
        if eq:
            por_equipo[(eq, letras)].append(suf)

    # 1) abreviatura de equipo: el tramo inicial que más se repite en las señales de ese equipo
    equipo: dict = {}
    for (eq, letras), sufs in por_equipo.items():
        cand = Counter(" ".join(s.split()[:n]) for s in sufs for n in (1, 2, 3))
        # el tramo más largo que aparece en al menos el 60% de las señales
        ok = [c for c, n in cand.items() if n >= 0.6 * len(sufs)]
        if ok:
            elegido = max(ok, key=lambda c: len(c.split()))
            if letras:
                equipo.setdefault(eq, {})[letras] = elegido
            else:
                equipo[eq] = elegido

    def abrev_eq(eq, letras):
        v = equipo.get(eq)
        return (v.get(letras) if isinstance(v, dict) else v) or ""

    # 2) frases (alarmas con equipo) y estados (señalizaciones con equipo)
    frase_c, estado_c = defaultdict(Counter), defaultdict(Counter)
    estados = {(e["s0"], e["s1"], e["s2"], e["s3"]): e["id"] for e in perfil.estados["estados"]}
    for eq, letras, frase, suf, est, obs in con_equipo:
        a = abrev_eq(eq, letras)
        if not a or not suf.startswith(a):
            continue
        resto = suf[len(a):].strip()
        eid = estados.get(tuple(str(x) if x is not None else None for x in est))
        if eid and eid.startswith(("sen_", "com_")):
            estado_c[eid][resto] += 1
        elif frase:
            frase_c[frase][resto] += 1

    # 3) señales sin equipo: la abreviatura completa de ese texto
    texto_c = defaultdict(Counter)
    for eq, letras, frase, suf, est, obs in sin_equipo:
        texto_c[re.sub(r"\d+", "#", dsc._norm(obs))][suf] += 1

    conflictos = []

    def elegir(nombre, tabla):
        out = {}
        for k, c in tabla.items():
            if len(c) > 1 and len({re.sub(r"\d", "#", v) for v in c}) == 1:
                conflictos.append(f"{nombre} «{k}» OMITIDA (depende de un número: {dict(c)})")
                continue
            (v, n), *otros = c.most_common()
            out[k] = v
            if otros:
                conflictos.append(f"{nombre} «{k}»: {dict(c)}")
        return out

    return {"equipo": equipo, "frase": elegir("frase", frase_c), "estado": elegir("estado", estado_c),
            "texto": elegir("texto", texto_c), "_conflictos": conflictos, "_filas": len(filas)}


def main(argv) -> int:
    if len(argv) < 2:
        print(__doc__)
        return 1
    perfil = perfiles.cargar_perfil(argv[2] if len(argv) > 2 else "EPEN")
    nuevo = extraer(argv[1], perfil)
    ruta = dsc.RAIZ_PLANTILLAS / perfil.id / "reglas_descripcion.json"
    actual = json.loads(ruta.read_text(encoding="utf-8"))
    for sec in ("equipo", "frase", "estado", "texto"):
        base = actual.get(sec, {})
        for k, v in nuevo[sec].items():
            if isinstance(v, dict) and isinstance(base.get(k), dict):
                base[k] = {**base[k], **v}
            else:
                base[k] = v
        actual[sec] = base
    actual["_origen_excel"] = f"{Path(argv[1]).name} ({nuevo['_filas']} filas); el Excel pisa a los lineamientos"
    actual["_conflictos"] = nuevo["_conflictos"]
    ruta.write_text(json.dumps(actual, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{ruta}\n  equipos: {nuevo['equipo']}\n  frases: {len(nuevo['frase'])}  estados: {nuevo['estado']}  "
          f"textos: {len(nuevo['texto'])}  conflictos: {len(nuevo['_conflictos'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
