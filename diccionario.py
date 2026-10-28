"""Diccionario de descripciones aprendido de un listado ya hecho (Excel).

Del Excel se toma, por cada fila, el texto de la señal (DESC. HMI sin el prefijo y OBSERVACIONES) y la
abreviatura que se usó en DESC. SCADA (sin el prefijo del campo). Con eso se sugiere la DESC. SCADA de
las señales nuevas que salen del plano. Los números (campo, equipo, tensión) se ignoran para que
"Seccionador 329SA03 Falta Tensión Auxiliar" y "Seccionador 229SA04 Falta Tensión Auxiliar" sean la
misma entrada. Nada se inventa: lo que no se parece a algo conocido queda para completar a mano.
"""
from __future__ import annotations

import difflib
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from modelo import Funcion, Proyecto, Senal
from perfiles import Perfil

SINONIMOS = {"secc": "seccionador", "ten": "tension", "int": "interruptor"}
CORTE_SIMILAR = 0.86


def normalizar(texto: str) -> str:
    t = unicodedata.normalize("NFKD", texto or "")
    t = "".join(c for c in t if not unicodedata.combining(c)).lower()
    t = re.sub(r"\d+", "#", t)
    t = re.sub(r"[^a-z#/ ]", " ", t)
    return " ".join(SINONIMOS.get(w, w) for w in t.split())


def _codigos(clave: str) -> set[str]:
    return {w for w in clave.split() if "#" in w and re.search(r"[a-z]", w)}


@dataclass
class Sugerencia:
    sufijo: str
    exacta: bool
    clave: str


@dataclass
class Diccionario:
    entradas: dict[str, Counter] = field(default_factory=lambda: defaultdict(Counter))
    filas_leidas: int = 0

    def __len__(self) -> int:
        return len(self.entradas)

    def agregar(self, textos: list[str], sufijo: str) -> None:
        for t in textos:
            k = normalizar(t)
            if k:
                self.entradas[k][sufijo] += 1

    def sugerir(self, texto: str) -> Optional[Sugerencia]:
        k = normalizar(texto)
        if not k:
            return None
        if k in self.entradas:
            return Sugerencia(self.entradas[k].most_common(1)[0][0], True, k)
        # parecidos: el código de equipo (#sa#, #sl#…) tiene que coincidir, porque cambia la abreviatura
        for c in difflib.get_close_matches(k, list(self.entradas), n=5, cutoff=CORTE_SIMILAR):
            if _codigos(c) == _codigos(k):
                return Sugerencia(self.entradas[c].most_common(1)[0][0], False, c)
        return None


def _prefijo(desc_hmi: str) -> str:
    return desc_hmi.split(" ", 1)[0] if desc_hmi else ""


def cargar_diccionario(ruta, perfil: Perfil, dic: Optional[Diccionario] = None) -> Diccionario:
    """Lee el listado (hoja y columnas del perfil) y suma sus descripciones al diccionario."""
    from openpyxl import load_workbook
    dic = dic if dic is not None else Diccionario()
    letra = {c["letra"]: c["clave"] for c in perfil.columnas}
    hojas = perfil.datos["hojas"]
    wb = load_workbook(Path(ruta), read_only=True, data_only=True)
    try:
        ws = wb[hojas["listado"]]
        for fila in ws.iter_rows(min_row=hojas["primera_fila_datos"]):
            d = {letra[c.column_letter]: c.value for c in fila
                 if getattr(c, "column_letter", None) in letra and c.value is not None}
            hmi, scada, obs = str(d.get("desc_hmi", "")).strip(), str(d.get("desc_scada", "")).strip(), \
                str(d.get("observaciones", "")).strip()
            if not hmi or not scada or str(d.get("tag_scada", "")).upper() == "N/A" or hmi.upper() == "RESERVA":
                continue
            prefijo = _prefijo(hmi)
            if not prefijo or not scada.startswith(prefijo):
                continue  # abreviatura con otro prefijo: no se puede separar con seguridad
            sufijo = scada[len(prefijo):].strip()
            if not sufijo:
                continue
            dic.filas_leidas += 1
            dic.agregar([hmi[len(prefijo):], obs], sufijo)
    finally:
        wb.close()
    return dic


@dataclass
class Informe:
    exactas: int = 0
    similares: list[tuple[str, str, str]] = field(default_factory=list)  # (texto del plano, parecido en el Excel, sufijo)
    sin_dato: list[str] = field(default_factory=list)


def aplicar(proyecto: Proyecto, dic: Diccionario, maximo: int = 20) -> Informe:
    """Completa DESC. SCADA de las señales que no la tienen. Marca el campo como fijo para que el
    generador no lo pise. Las sugerencias 'similares' se aplican igual pero quedan listadas para revisar."""
    inf = Informe()
    for nivel in proyecto.niveles:
        for campo in nivel.campos:
            prefijo = campo.prefijo or f"{nivel.codigo}{campo.numero}"
            propuestas: list[tuple[Senal, str, Sugerencia]] = []
            usadas = Counter(s.desc_scada for s in campo.senales if s.desc_scada)
            for s in campo.senales:
                if s.segunda_fila or s.es_reserva or s.funcion == Funcion.RESERVA or s.desc_scada \
                        or "desc_scada" in s.fijos or not s.observaciones:
                    continue
                sug = dic.sugerir(s.observaciones)
                if not sug:
                    inf.sin_dato.append(s.observaciones)
                    continue
                desc = _unir(prefijo, sug.sufijo, maximo)
                if not desc:
                    inf.sin_dato.append(f"{s.observaciones} (la abreviatura '{sug.sufijo}' no entra en {maximo})")
                    continue
                propuestas.append((s, desc, sug))
                usadas[desc] += 1
            for s, desc, sug in propuestas:
                if usadas[desc] > 1:  # dos señales distintas del campo no pueden llamarse igual
                    inf.sin_dato.append(f"{s.observaciones} (la sugerencia '{desc}' se repetiría en el campo)")
                    continue
                s.desc_scada = desc
                s.fijos.append("desc_scada")
                if sug.exacta:
                    inf.exactas += 1
                else:
                    inf.similares.append((s.observaciones, sug.clave, desc))
    return inf


def _unir(prefijo: str, sufijo: str, maximo: int) -> str:
    for desc in (f"{prefijo} {sufijo}", f"{prefijo}{sufijo}"):  # el Excel a veces pega el prefijo
        if len(desc) <= maximo:
            return desc
    return ""
