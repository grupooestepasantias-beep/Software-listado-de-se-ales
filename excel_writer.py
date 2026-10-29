"""Escritor del listado de señales a .xlsx.

Hojas, columnas activas, letras y orden salen del perfil; anchos, estilos, merges de
encabezado, filtros, impresión y posición de los datos de la carátula salen de la
plantilla (plantillas/<id_perfil>/<id_perfil>_plantilla.json, extraída del ejemplo con
herramientas/extraer_plantilla.py). Aquí no hay reglas de ningún cliente.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Optional, Union

from openpyxl import Workbook
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.properties import PageSetupProperties

from modelo import Proyecto, Senal
from perfiles import Perfil

RAIZ_PLANTILLAS = Path(__file__).parent / "plantillas"


class PlantillaNoEncontrada(FileNotFoundError):
    pass


class DocumentoInvalido(ValueError):
    pass


def cargar_plantilla(perfil: Perfil) -> dict:
    """Usa perfil.datos["plantilla"] si existe; si no, plantillas/<id>/."""
    rel = perfil.datos.get("plantilla")
    carpeta = Path(__file__).parent / rel if rel else RAIZ_PLANTILLAS / perfil.id
    ruta = carpeta / f"{perfil.id}_plantilla.json"
    if not ruta.is_file():
        raise PlantillaNoEncontrada(f"no existe la plantilla {ruta}")
    return json.loads(ruta.read_text(encoding="utf-8"))


# ---------------------------------------------------------------- estilos
def _color(spec) -> Optional[Color]:
    if spec is None:
        return None
    if isinstance(spec, str):
        return Color(rgb=spec)
    if "theme" in spec:
        return Color(theme=spec["theme"], tint=spec.get("tint", 0.0))
    if "indexed" in spec:
        return Color(indexed=spec["indexed"])
    return None


def _lado(spec) -> Side:
    if not spec:
        return Side()
    return Side(style=spec["style"], color=_color(spec.get("color")))


def aplicar_estilo(cell, e: dict) -> None:
    f = e["font"]
    cell.font = Font(name=f["name"], sz=f["sz"], b=f["b"], i=f["i"], u=f["u"],
                     color=_color(f["color"]))
    if e["fill"]:
        cell.fill = PatternFill(fill_type=e["fill"]["type"], fgColor=_color(e["fill"]["fg"]))
    b = e["border"]
    cell.border = Border(left=_lado(b["left"]), right=_lado(b["right"]),
                         top=_lado(b["top"]), bottom=_lado(b["bottom"]))
    a = e["align"]
    cell.alignment = Alignment(horizontal=a["h"], vertical=a["v"], wrap_text=a["wrap"],
                               textRotation=a["rot"], indent=a["indent"])
    cell.number_format = e["fmt"]


def _anchos(ws, columnas: dict) -> None:
    for letra, c in columnas.items():
        ancho = c["ancho"] if isinstance(c, dict) else c
        if ancho:
            ws.column_dimensions[letra].width = ancho
        if isinstance(c, dict) and c.get("oculta"):
            ws.column_dimensions[letra].hidden = True


def _con_fila_final(ref: str, ultima: int) -> str:
    """'A5:W584' -> 'A5:W{ultima}'; también sirve para áreas con comillas/$."""
    return re.sub(r"(\$?[A-Z]+\$?)\d+(?!.*\d)", lambda m: f"{m.group(1)}{ultima}", ref)


# ---------------------------------------------------------------- listado
def _valor(clave: str, s: Senal, item: int, nivel, campo) -> object:
    if clave == "item":
        return item
    if clave == "nivel":
        return nivel.nombre
    if clave == "descrip":
        return campo.descripcion
    if clave == "controlador":
        return s.controlador or campo.controlador
    if clave == "funcion":
        return s.funcion.value
    if clave == "tipo_punto":
        return s.tipo_punto.value if s.tipo_punto else None
    return getattr(s, clave, None)


def _escribir_listado(ws, proyecto: Proyecto, perfil: Perfil, t: dict, omitir: frozenset = frozenset(),
                      hasta: Optional[str] = None) -> int:
    hojas = perfil.datos["hojas"]
    primera = hojas["primera_fila_datos"]
    enc = hojas["fila_encabezado"]
    limite = column_index_from_string(hasta) if hasta else 10 ** 6  # columnas más allá de `hasta` no se escriben
    dentro = lambda letra: column_index_from_string(letra) <= limite  # noqa: E731
    activas = [c for c in perfil.columnas_activas if dentro(c["letra"])]
    letra_de = {c["clave"]: c["letra"] for c in perfil.columnas}

    _anchos(ws, {k: v for k, v in t["columnas"].items() if dentro(k)})
    for r, h in t["alturas_filas"].items():
        ws.row_dimensions[int(r)].height = h
    for r, fila in t["encabezado"].items():
        for letra, cel in fila.items():
            if dentro(letra):
                c = ws[f"{letra}{r}"]
                c.value = cel["v"]
                aplicar_estilo(c, cel)
    for m in t["merges_encabezado"]:
        if CellRange(m).max_col <= limite:
            ws.merge_cells(m)

    estilo_dato = t["estilo_filas_datos"][str(primera)]
    alto = t["alturas_filas"].get(str(primera))
    fila = primera
    item = proyecto.item_inicial
    # (columna -> [(desde, hasta)]) de merges verticales, solo en columnas activas
    merges: list[tuple[str, int, int]] = []

    def _cerrar(letra: str, desde: int, hasta: int) -> None:
        if hasta > desde:
            merges.append((letra, desde, hasta))

    for nivel in proyecto.niveles:
        n_desde = fila
        for campo in nivel.campos:
            c_desde = fila
            f_desde, f_func = fila, None
            for s in campo.senales:
                if f_func is not None and s.funcion != f_func:
                    _cerrar(letra_de["funcion"], f_desde, fila - 1)
                    f_desde = fila
                f_func = s.funcion
                for col in activas:
                    cel = ws[f"{col['letra']}{fila}"]
                    if col["clave"] not in omitir:  # columnas de una etapa posterior (ej. SCADA): quedan vacías
                        v = _valor(col["clave"], s, item, nivel, campo)
                        cel.value = None if v == "" else v
                    aplicar_estilo(cel, estilo_dato[col["letra"]])
                if alto:
                    ws.row_dimensions[fila].height = alto
                item += 1
                fila += 1
            if f_func is not None:
                _cerrar(letra_de["funcion"], f_desde, fila - 1)
            _cerrar(letra_de["descrip"], c_desde, fila - 1)
        _cerrar(letra_de["nivel"], n_desde, fila - 1)

    activas_letras = {c["letra"] for c in activas}
    for letra, d, h in merges:
        if letra in activas_letras:
            ws.merge_cells(f"{letra}{d}:{letra}{h}")

    ultima = max(fila - 1, primera)  # sin señales el rango queda en la primera fila de datos
    if t["autofiltro"]:
        ref = _con_fila_final(t["autofiltro"], ultima)
        ws.auto_filter.ref = re.sub(r"[A-Z]+(\d+)$", rf"{hasta}\1", ref) if hasta else ref
    if t["area_impresion"]:
        area = _con_fila_final(t["area_impresion"].split("!")[-1].replace("$", ""), ultima)
        ws.print_area = re.sub(r"[A-Z]+(\d+)$", rf"{hasta}\1", area) if hasta else area
    if t["titulos_repetir_filas"]:
        ws.print_title_rows = t["titulos_repetir_filas"].replace("$", "")
    _pagina(ws, t)

    # valores duplicados en DESC. SCADA (misma regla visual que el ejemplo)
    if t["formato_condicional"] and "desc_scada" in letra_de and fila > primera and dentro(letra_de["desc_scada"]):
        regla = t["formato_condicional"][0]
        dxf = DifferentialStyle(fill=PatternFill(bgColor=_color(regla["dxf_fill"])))
        L = letra_de["desc_scada"]
        ws.conditional_formatting.add(f"{L}{primera}:{L}{ultima}",
                                      Rule(type="duplicateValues", dxf=dxf))
    return ultima


def _pagina(ws, t: dict) -> None:
    p = t["pagina"]
    ws.page_setup.orientation = p["orientacion"]
    ws.page_setup.paperSize = p["papel"]
    if p.get("escala"):
        ws.page_setup.scale = p["escala"]
    if p.get("fit_to_page"):
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToHeight = p.get("ajustar_alto", 0)
        ws.page_setup.fitToWidth = p.get("ajustar_ancho") or 1
    for k, v in t["margenes"].items():
        setattr(ws.page_margins, k, v)


# ---------------------------------------------------------------- carátula
def _coords_dinamicas(campos: dict) -> set[str]:
    out: set[str] = set()
    r0, r1 = campos["referencias"]["filas"]
    for r in range(r0, r1 + 1):
        out.add(f'{campos["referencias"]["col_codigo"]}{r}')
        out.add(f'{campos["referencias"]["col_titulo"]}{r}')
    v0, v1 = campos["revisiones"]["filas"]
    for r in range(v0, v1 + 1):
        out.update(f"{col}{r}" for col in campos["revisiones"]["columnas"].values())
    out.update(c["celda"] for c in campos["simples"].values())
    out.update(campos["documento"]["celdas"])
    out.add(campos["hoja"]["celda"])
    return out


def _escribir_caratula(ws, proyecto: Proyecto, perfil: Perfil, t: dict, campos: dict) -> None:
    car = proyecto.caratula
    dinamicas = _coords_dinamicas(campos)

    _anchos(ws, t["columnas"])
    for r, h in t["alturas_filas"].items():
        ws.row_dimensions[int(r)].height = h
    for coord, cel in t["celdas"].items():
        c = ws[coord]
        c.value = None if coord in dinamicas else cel["v"]
        aplicar_estilo(c, cel)
    for m in t["merges"]:
        ws.merge_cells(m)

    # REFERENCIAS
    ref = campos["referencias"]
    r0, r1 = ref["filas"]
    if len(car.referencias) > r1 - r0 + 1:
        raise ValueError(f"la carátula admite {r1 - r0 + 1} referencias, hay {len(car.referencias)}")
    for i, rf in enumerate(car.referencias):
        ws[f'{ref["col_codigo"]}{r0 + i}'] = rf.codigo
        ws[f'{ref["col_titulo"]}{r0 + i}'] = rf.descripcion

    # REVISIONES: la más reciente arriba, se llenan de abajo hacia arriba (lista: de la más
    # antigua a la más nueva)
    v0, v1 = campos["revisiones"]["filas"]
    if len(car.revisiones) > v1 - v0 + 1:
        raise ValueError(f"la carátula admite {v1 - v0 + 1} revisiones, hay {len(car.revisiones)}")
    for i, rv in enumerate(car.revisiones):
        for atributo, col in campos["revisiones"]["columnas"].items():
            ws[f"{col}{v1 - i}"] = getattr(rv, atributo)

    # Campos simples (con formato opcional: "   Lugar: {}")
    for nombre, cfg in campos["simples"].items():
        valor = getattr(car, nombre if nombre != "dibujo" else "dibujo")
        fmt = cfg.get("formato")
        ws[cfg["celda"]] = fmt.format(valor) if fmt else valor

    # Número de documento: un carácter por celda, sin guiones. Llega armado (Caratula.numero_documento).
    celdas = campos["documento"]["celdas"]
    caracteres = car.numero_documento.replace("-", "")
    if len(caracteres) > len(celdas):
        raise DocumentoInvalido(
            f"el número de documento '{car.numero_documento}' tiene {len(caracteres)} caracteres "
            f"(sin guiones) y la plantilla admite hasta {len(celdas)}")
    for i, coord in enumerate(celdas):
        ws[coord] = caracteres[i] if i < len(caracteres) else None

    # Hoja X de Y: el total de hojas es del perfil
    h = campos["hoja"]
    ws[h["celda"]] = h["formato"].format(n=1, total=perfil.datos["hojas"]["total"])

    if t["area_impresion"]:
        ws.print_area = t["area_impresion"].split("!")[-1].replace("$", "")
    _pagina_caratula(ws, t)
    _ajustar_caratula(ws)
    _agregar_logos(ws, car)


def _pagina_caratula(ws, t: dict) -> None:
    p = t["pagina"]
    ws.page_setup.orientation = p["orientacion"]
    ws.page_setup.paperSize = p["papel"]
    if p.get("escala"):
        ws.page_setup.scale = p["escala"]
    for k, v in t["margenes"].items():
        setattr(ws.page_margins, k, v)


def _ajustar_caratula(ws) -> None:
    """La carátula se ve e imprime en UNA hoja A4 vertical, con las alturas/anchos por defecto del ETAP."""
    from openpyxl.utils import get_column_letter
    ws.sheet_format.defaultRowHeight = 12.75
    ws.sheet_format.defaultColWidth = 11.42578125
    # En el ETAP las columnas C..AD comparten el ancho de la C (grupo C:AD); la plantilla solo guardó la C.
    ancho_c = ws.column_dimensions["C"].width
    for i in range(4, 31):
        ws.column_dimensions[get_column_letter(i)].width = ancho_c
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_options.verticalCentered = True


# Recuadros de los logos en la carátula del ETAP: (col, desplazamiento EMU, fila, desplazamiento EMU; desde 0).
# La imagen elegida se escala para entrar en el recuadro SIN deformarse, anclada arriba a la izquierda.
CAJAS_LOGO = {
    "logo_cliente": {"desde": (0, 180975, 31, 47625), "tam": (752475, 742950)},
    "logo_proyectista": {"desde": (0, 22412, 45, 27945), "hasta": (1, 907678, 48, 123264)},
}
_EMU_PX, _EMU_PT = 9525, 12700
FORMATOS_IMAGEN = (".png", ".jpg", ".jpeg", ".gif", ".bmp")


def _tam_caja(ws, caja: dict) -> tuple[int, int]:
    """Ancho y alto del recuadro en EMU (las columnas y filas ya tienen su ancho/alto final)."""
    from openpyxl.utils import get_column_letter
    if "tam" in caja:
        return caja["tam"]
    c0, co0, r0, ro0 = caja["desde"]
    c1, co1, r1, ro1 = caja["hasta"]

    def ancho(i):
        w = ws.column_dimensions[get_column_letter(i + 1)].width or ws.sheet_format.defaultColWidth or 8.43
        return int((w * 7 + 5) * _EMU_PX)

    def alto(i):
        return int((ws.row_dimensions[i + 1].height or ws.sheet_format.defaultRowHeight or 15) * _EMU_PT)

    return (sum(ancho(i) for i in range(c0, c1)) - co0 + co1, sum(alto(i) for i in range(r0, r1)) - ro0 + ro1)


def validar_imagen(ruta: str, nombre: str) -> list[str]:
    """Problemas de un logo ('' = sin logo, es válido)."""
    if not ruta:
        return []
    p = Path(ruta)
    if not p.is_file():
        return [f"{nombre}: no se encuentra el archivo {ruta}"]
    if p.suffix.lower() not in FORMATOS_IMAGEN:
        return [f"{nombre}: formato no admitido ({p.suffix}); usá {', '.join(FORMATOS_IMAGEN)}"]
    try:
        from PIL import Image as PILImage
        with PILImage.open(p) as im:
            im.verify()
    except Exception:
        return [f"{nombre}: el archivo no es una imagen válida"]
    return []


def _agregar_logos(ws, car) -> None:
    """Logos elegidos por el usuario (Caratula.logo_cliente / logo_proyectista; '' = sin logo)."""
    import io

    from openpyxl.drawing.image import Image
    from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from PIL import Image as PILImage
    for clave, caja in CAJAS_LOGO.items():
        ruta = getattr(car, clave, "") or ""
        if not ruta:
            continue
        problemas = validar_imagen(ruta, clave)
        if problemas:
            raise ValueError(problemas[0])
        with PILImage.open(ruta) as im:  # se achica si es enorme (máx. 500 px de lado) y va como PNG
            im = im.convert("RGBA") if im.mode in ("P", "LA", "RGBA") else im.convert("RGB")
            im.thumbnail((500, 500))
            w, h = im.size
            buf = io.BytesIO()
            im.save(buf, format="PNG")
            buf.seek(0)
        cw, ch = _tam_caja(ws, caja)
        escala = min(cw / w, ch / h)
        c, co, r, ro = caja["desde"]
        img = Image(buf)
        img.anchor = OneCellAnchor(_from=AnchorMarker(col=c, colOff=co, row=r, rowOff=ro),
                                   ext=XDRPositiveSize2D(int(w * escala), int(h * escala)))
        ws.add_image(img)


# ---------------------------------------------------------------- API
def escribir_xlsx(proyecto: Proyecto, perfil: Perfil, ruta: Union[str, Path],
                  plantilla: Optional[dict] = None, incluir_caratula: bool = True,
                  omitir_claves=(), hasta_columna: Optional[str] = None) -> Path:
    """Genera el .xlsx (hoja de carátula + hoja del listado) y devuelve la ruta escrita.

    Con incluir_caratula=False el libro lleva solo la hoja del listado. `omitir_claves` = claves de
    columnas (perfil) que no se escriben (quedan vacías, con su formato). `hasta_columna` = última letra
    que se escribe (ej. "G": el listado sale solo con las columnas A-G)."""
    t = plantilla or cargar_plantilla(perfil)
    hojas = perfil.datos["hojas"]
    wb = Workbook()
    if incluir_caratula:
        ws_car = wb.active
        ws_car.title = hojas["caratula"]
        _escribir_caratula(ws_car, proyecto, perfil, t["caratula"], t["caratula_campos"])
        ws_lis = wb.create_sheet(hojas["listado"])
    else:
        ws_lis = wb.active
        ws_lis.title = hojas["listado"]
    _escribir_listado(ws_lis, proyecto, perfil, t["listado"], frozenset(omitir_claves), hasta_columna)
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta
