import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import excel_writer
import perfiles
from modelo import Campo, Caratula, Funcion, Nivel, Proyecto, Referencia, Revision, Senal, TipoPunto

try:
    from openpyxl import load_workbook
except ImportError:  # el módulo de datos no la necesita, pero estos tests sí
    load_workbook = None

EJEMPLO = Path(r"\\192.168.15.120\PasantiasMM\Sofware listado\Ejemplos\ETAP-0722-ET-CT-MC-LI-6209-1-Lista señales - CAO.xlsx")
PRIMERA, ULTIMA = 6, 84  # tramo del ejemplo que se reproduce (ALARMAS + SEÑALIZACION del campo 303)
INACTIVAS = {"Q", "R", "U", "V"}  # columnas fuera de alcance (perfil: activa false)


def _norm(v):
    return None if v in (None, "") else v


def _proyecto_desde_ejemplo(ws, perfil) -> Proyecto:
    """Arma el Proyecto de prueba con filas reales del ejemplo (datos fijos)."""
    letra = {c["clave"]: c["letra"] for c in perfil.columnas}
    senales, funcion = [], None
    for r in range(PRIMERA, ULTIMA + 1):
        g = lambda k: ws[f"{letra[k]}{r}"].value
        funcion = Funcion(g("funcion")) if g("funcion") else funcion
        tp = g("tipo_punto")
        senales.append(Senal(
            funcion=funcion, tipo_punto=TipoPunto(tp) if tp else None,
            controlador=g("controlador") or "", tag_hmi=g("tag_hmi") or "",
            desc_hmi=g("desc_hmi") or "", tag_scada=g("tag_scada") or "",
            offset=g("offset"), desc_scada=g("desc_scada") or "",
            s0=g("s0"), s1=g("s1"), s2=g("s2"), s3=g("s3"), unidad=g("unidad"),
            placa=g("placa"), entrada=g("entrada"), observaciones=g("observaciones") or ""))
    campo = Campo(numero="03", descripcion=ws[f"{letra['descrip']}{PRIMERA}"].value, senales=senales)
    car = Caratula(
        cliente="ENTE PROVINCIAL DE ENERGÍA DEL NEUQUÉN", obra="LAT 132 kV LOMA CAMPANA-AGUADA PICHANA",
        lugar="LOMA CAMPANA", departamento="AÑELO", titulo="ETAP - LISTADO DE SEÑALES",
        numero_documento="ETAP-0722-ET-CT-MC-LI-6209", proyecto="D.Vera", dibujo="G. García",
        aprobo="M. Escalante", archivo_magnetico="ETAP-0722-ET-CT-MC-LI-6209", escala="N/A",
        revision="1", fecha="2023-05-11",
        referencias=[Referencia("ETAP-0722-ET-CT-PL-EF-6106",
                                "ET AP - TC3 03 - TABLERO DE CONTROL DE LÍNEAS CAMPO 03_132Kv - ESQUEMA FUNCIONAL")],
        revisiones=[Revision("0", "21/04/2025", "APC", "DV", "GG", "ME"),
                    Revision("1", "08/03/2026", "CONFORME A OBRA", "DV", "GG", "ME")])
    return Proyecto(perfil_id=perfil.id, caratula=car, item_inicial=ws[f"{letra['item']}{PRIMERA}"].value,
                    niveles=[Nivel("132KV", "3", [campo])])


@unittest.skipUnless(load_workbook and EJEMPLO.exists(), "falta openpyxl o el ejemplo ETAP")
class TestExcelWriter(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.perfil = perfiles.cargar_perfil("EPEN")
        cls.orig = load_workbook(EJEMPLO)
        cls.proyecto = _proyecto_desde_ejemplo(cls.orig[cls.perfil.datos["hojas"]["listado"]], cls.perfil)
        cls.tmp = tempfile.TemporaryDirectory()
        cls.ruta = excel_writer.escribir_xlsx(cls.proyecto, cls.perfil, Path(cls.tmp.name) / "salida.xlsx")
        cls.gen = load_workbook(cls.ruta)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_hojas_del_perfil(self):
        h = self.perfil.datos["hojas"]
        self.assertEqual(self.gen.sheetnames, [h["caratula"], h["listado"]])

    def test_filas_iguales_al_ejemplo_en_columnas_activas(self):
        o, g = self.orig["EPEN"], self.gen["EPEN"]
        difs = []
        for c in self.perfil.columnas_activas:
            for r in range(PRIMERA, ULTIMA + 1):
                a, b = _norm(o[f"{c['letra']}{r}"].value), _norm(g[f"{c['letra']}{r}"].value)
                if a != b:
                    difs.append((f"{c['letra']}{r}", a, b))
        self.assertEqual(difs, [], difs[:5])

    def test_columnas_inactivas_vacias_y_sin_merges(self):
        g = self.gen["EPEN"]
        for L in INACTIVAS:
            for r in range(PRIMERA, ULTIMA + 1):
                self.assertIsNone(g[f"{L}{r}"].value)
        for m in g.merged_cells.ranges:
            if m.min_row >= PRIMERA:
                self.assertNotIn(m.min_col, (17, 18, 21, 22), str(m))  # Q, R, U, V

    def test_merges_de_datos_B_C_E(self):
        o, g = self.orig["EPEN"], self.gen["EPEN"]

        def merges(ws):
            return {str(m) for m in ws.merged_cells.ranges
                    if m.min_row >= PRIMERA and m.min_col in (2, 3, 5)}

        esperado = {f"{m[0]}{PRIMERA}:{m[0]}{ULTIMA}" for m in ("B", "C")}
        esperado |= {str(m) for m in o.merged_cells.ranges
                     if m.min_col == 5 and m.min_row >= PRIMERA and m.max_row <= ULTIMA}
        self.assertEqual(merges(g), esperado)

    def test_encabezado_y_hoja(self):
        o, g = self.orig["EPEN"], self.gen["EPEN"]
        for r in range(1, 6):
            for c in range(1, 24):
                self.assertEqual(o.cell(r, c).value, g.cell(r, c).value, f"fila {r} col {c}")
        self.assertEqual(g.auto_filter.ref, f"A5:W{ULTIMA}")
        self.assertEqual(g.print_title_rows.replace("$", ""), "1:5")
        for L, dim in o.column_dimensions.items():
            if dim.width:
                self.assertAlmostEqual(g.column_dimensions[L].width, dim.width, places=3)

    def test_caratula(self):
        o, g = self.orig["Caratula"], self.gen["Caratula"]
        for coord in ("B2", "C29", "I29", "C31", "C34", "A38", "H38", "A41", "C46", "C48", "C50",
                      "J47", "R46", "X46", "T49", "A27", "I27", "A28", "I28", "B3", "G3", "N44"):
            exp = o[coord].value
            if coord == "T49":
                exp = f"1 de {self.perfil.datos['hojas']['total']}"
            if coord == "A27":
                exp = "1"
            if coord == "A28":
                exp = "0"
            if coord == "X46":
                exp = "1"
            self.assertEqual(_norm(g[coord].value), _norm(exp), coord)
        self.assertEqual(sorted(map(str, g.merged_cells.ranges)), sorted(map(str, o.merged_cells.ranges)))

    def test_numero_documento_un_caracter_por_celda(self):
        g = self.gen["Caratula"]
        celdas = ["D44", "E44", "F44", "G44", "I44", "J44", "K44", "L44", "O44", "P44",
                  "R44", "S44", "U44", "V44", "X44", "Y44", "AA44", "AB44", "AC44", "AD44"]
        self.assertEqual("".join(str(g[c].value) for c in celdas), "ETAP0722ETCTMCLI6209")

    def test_documento_demasiado_largo(self):
        p = Proyecto(perfil_id="EPEN", caratula=Caratula(numero_documento="ETAP-0722-ET-CT-MC-LI-62099"))
        with self.assertRaises(excel_writer.DocumentoInvalido):
            excel_writer.escribir_xlsx(p, self.perfil, Path(self.tmp.name) / "malo.xlsx")

    def test_documento_mas_corto_deja_celdas_vacias(self):
        p = Proyecto(perfil_id="EPEN", caratula=Caratula(numero_documento="AB-12"))
        ruta = excel_writer.escribir_xlsx(p, self.perfil, Path(self.tmp.name) / "corto.xlsx")
        g = load_workbook(ruta)["Caratula"]
        self.assertEqual([g[c].value for c in ("D44", "E44", "F44", "G44", "I44")], ["A", "B", "1", "2", None])
        self.assertIsNone(g["AD44"].value)


if __name__ == "__main__":
    unittest.main()
