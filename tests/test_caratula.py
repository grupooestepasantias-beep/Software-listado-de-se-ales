"""Carátula integrada: logos, una hoja A4, anchos, datos de la pestaña y comparación celda por celda con el ETAP."""
import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import excel_writer  # noqa: E402
import perfiles  # noqa: E402
from app import caratula_ui as cu  # noqa: E402
from modelo import Caratula, Proyecto, Referencia, Revision  # noqa: E402

ETAP = Path(__file__).resolve().parent.parent.parent / "Ejemplos" / "ETAP-0722-ET-CT-MC-LI-6209-1-Lista señales - CAO.xlsx"
LOGO1, LOGO2 = cu.LOGO_CLIENTE_ETAP, cu.LOGO_PROYECTISTA_ETAP


def caratula(**kw) -> Caratula:
    base = dict(cliente="ENTE PROVINCIAL DE ENERGÍA DEL NEUQUÉN", obra="LAT 132 kV LOMA CAMPANA-AGUADA PICHANA",
                lugar="LOMA CAMPANA", departamento="AÑELO", titulo="ETAP - LISTADO DE SEÑALES",
                numero_documento="ETAP-0722-ET-CT-MC-LI-6209", proyecto="D.Vera", dibujo="G. García",
                aprobo="M. Escalante", archivo_magnetico="ETAP-0722-ET-CT-MC-LI-6209", revision="1", fecha="11/05/2023",
                referencias=[Referencia("ETAP-0722-ET-CT-PL-EF-6106", "ET AP - TC3 03")],
                revisiones=[Revision("0", "21/04/2025", "APC", "DV", "GG", "ME"),
                            Revision("1", "08/03/2026", "CONFORME A OBRA", "DV", "GG", "ME")],
                logo_cliente=str(LOGO1), logo_proyectista=str(LOGO2))
    base.update(kw)
    return Caratula(**base)


def escribir(car: Caratula, carpeta: str, nombre: str = "c.xlsx") -> Path:
    perfil = perfiles.cargar_perfil("EPEN")
    return excel_writer.escribir_xlsx(Proyecto("EPEN", caratula=car), perfil, Path(carpeta) / nombre)


class TestExcel(unittest.TestCase):
    def test_hojas_y_datos(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as t:
            wb = openpyxl.load_workbook(escribir(caratula(), t))
            self.assertEqual(wb.sheetnames, ["Caratula", "EPEN"])
            ws = wb["Caratula"]
            self.assertEqual(ws["C31"].value, "ENTE PROVINCIAL DE ENERGÍA DEL NEUQUÉN")
            self.assertEqual(ws["A38"].value, "   Lugar: LOMA CAMPANA")
            self.assertEqual(ws["T49"].value, "1 de 2")
            self.assertEqual((ws["A28"].value, ws["A27"].value), ("0", "1"))  # la más nueva arriba

    def test_logos_sin_sello(self):
        with tempfile.TemporaryDirectory() as t:
            medios = {i.filename: i.file_size for i in zipfile.ZipFile(escribir(caratula(), t)).infolist()
                      if i.filename.startswith("xl/media/")}
            self.assertEqual(len(medios), 2)                         # logo del cliente + logo del proyectista
            self.assertTrue(all(tam < 20000 for tam in medios.values()))  # el sello "CONFORME A OBRA" pesa ~137 KB
            sin = escribir(caratula(logo_cliente="", logo_proyectista=""), t, "s.xlsx")
            self.assertFalse([n for n in zipfile.ZipFile(sin).namelist() if n.startswith("xl/media/")])

    def test_logos_propios_sin_deformarse(self):
        import openpyxl
        from PIL import Image
        with tempfile.TemporaryDirectory() as t:
            ancho, grande = Path(t) / "ancho.png", Path(t) / "grande.png"
            Image.new("RGB", (400, 100), (200, 0, 0)).save(ancho)
            Image.new("RGB", (3000, 3000), (0, 0, 200)).save(grande)
            ruta = escribir(caratula(logo_cliente=str(ancho), logo_proyectista=str(grande)), t)
            medios = [i for i in zipfile.ZipFile(ruta).infolist() if i.filename.startswith("xl/media/")]
            self.assertTrue(all(i.file_size < 20000 for i in medios))     # la de 3000 px se achicó
            imgs = openpyxl.load_workbook(ruta)["Caratula"]._images
            por_fila = {im.anchor._from.row: im.anchor.ext for im in imgs}
            c1, c2 = por_fila[31], por_fila[45]
            self.assertAlmostEqual(c1.width / c1.height, 4.0, places=1)
            self.assertAlmostEqual(c2.width / c2.height, 1.0, places=1)
            self.assertLessEqual(c1.width, 752475 + 1)

    def test_una_pagina_y_anchos(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as t:
            ws = openpyxl.load_workbook(escribir(caratula(), t))["Caratula"]
            self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage)
            self.assertEqual((ws.page_setup.fitToWidth, ws.page_setup.fitToHeight), (1, 1))
            self.assertEqual(ws.column_dimensions["AD"].width, ws.column_dimensions["C"].width)  # grupo C:AD

    def test_logo_invalido_da_error_claro(self):
        with tempfile.TemporaryDirectory() as t:
            with self.assertRaises(ValueError):
                escribir(caratula(logo_cliente=str(Path(t) / "no_existe.png")), t)

    def test_validar_imagen(self):
        with tempfile.TemporaryDirectory() as t:
            falso, texto = Path(t) / "falso.png", Path(t) / "logo.txt"
            falso.write_text("no es una imagen")
            texto.write_text("x")
            self.assertEqual(excel_writer.validar_imagen("", "Logo"), [])
            self.assertIn("no se encuentra", excel_writer.validar_imagen(str(Path(t) / "x.png"), "Logo")[0])
            self.assertIn("no es una imagen válida", excel_writer.validar_imagen(str(falso), "Logo")[0])
            self.assertIn("formato no admitido", excel_writer.validar_imagen(str(texto), "Logo")[0])
            self.assertEqual(excel_writer.validar_imagen(str(LOGO1), "Logo"), [])


class TestPestana(unittest.TestCase):
    def test_tablas_como_texto(self):
        refs = [Referencia("A-1", "Plano uno"), Referencia("A-2", "Plano dos")]
        self.assertEqual(cu.texto_a_referencias(cu.referencias_a_texto(refs)), refs)
        revs = caratula().revisiones
        self.assertEqual(cu.texto_a_revisiones(cu.revisiones_a_texto(revs)), revs)
        self.assertEqual(cu.texto_a_revisiones("A | 01/01/2026 | PARA APROBACIÓN")[0].aprobador, "")

    def test_tiene_datos(self):
        self.assertFalse(cu.tiene_datos(Caratula()))                      # solo la escala N/A de fábrica
        self.assertFalse(cu.tiene_datos(cu.caratula_nueva()))             # los logos preseleccionados no cuentan
        self.assertTrue(cu.tiene_datos(Caratula(obra="X")))
        self.assertTrue(cu.tiene_datos(Caratula(referencias=[Referencia("A", "b")])))

    def test_limites_y_problemas(self):
        lim = cu.limites_de(perfiles.cargar_perfil("EPEN"))
        self.assertEqual(lim, {"referencias": 14, "revisiones": 4, "documento": 20})
        self.assertEqual(cu.problemas(caratula(), lim), [])
        mal = caratula(numero_documento="A" * 21, referencias=[Referencia(f"R{i}", "x") for i in range(15)],
                       revisiones=[Revision("", "x")], logo_cliente="no_existe.png")
        m = "\n".join(cu.problemas(mal, lim))
        for esperado in ("hasta 20 caracteres", "15 referencias", "falta la letra", "no se encuentra el archivo"):
            self.assertIn(esperado, m)

    def test_caratula_nueva_trae_los_logos_del_etap(self):
        c = cu.caratula_nueva()
        self.assertEqual((c.logo_cliente, c.logo_proyectista), (str(LOGO1), str(LOGO2)))

    def test_json_con_logos(self):
        p = Proyecto("EPEN", caratula=caratula())
        self.assertEqual(Proyecto.from_dict(p.to_dict()), p)
        viejo = p.to_dict()
        del viejo["caratula"]["logo_cliente"], viejo["caratula"]["logo_proyectista"]
        self.assertEqual(Proyecto.from_dict(viejo).caratula.logo_cliente, "")     # proyectos viejos: sin logo


@unittest.skipUnless(ETAP.exists(), "ejemplo ETAP no disponible")
class TestContraEtap(unittest.TestCase):
    """Reconstruye los datos desde la hoja Caratula del ETAP, los escribe y compara celda por celda."""

    def test_identica_al_ejemplo(self):
        import openpyxl
        orig = openpyxl.load_workbook(ETAP, data_only=True)["Caratula"]   # el nº de documento son fórmulas MID()
        c = excel_writer.cargar_plantilla(perfiles.cargar_perfil("EPEN"))["caratula_campos"]

        def txt(v):  # ojo: el 0 es un valor válido (dígito del nº de documento, letra de revisión)
            return "" if v is None else str(v)

        d = {}
        for nombre, cfg in c["simples"].items():
            v = txt(orig[cfg["celda"]].value)
            if cfg.get("formato"):
                prefijo = cfg["formato"].split("{}")[0]
                v = v[len(prefijo):] if v.startswith(prefijo) else ("" if v.rstrip() == prefijo.rstrip() else v)
            d[nombre] = v
        d["numero_documento"] = "".join(txt(orig[x].value) for x in c["documento"]["celdas"])
        r0, r1 = c["referencias"]["filas"]
        cc, ct = c["referencias"]["col_codigo"], c["referencias"]["col_titulo"]
        d["referencias"] = [Referencia(txt(orig[f"{cc}{r}"].value), txt(orig[f"{ct}{r}"].value))
                            for r in range(r0, r1 + 1) if orig[f"{cc}{r}"].value]
        v0, v1 = c["revisiones"]["filas"]
        d["revisiones"] = []
        for r in range(v1, v0 - 1, -1):  # de abajo (más antigua) hacia arriba
            fila = {k: orig[f"{col}{r}"].value for k, col in c["revisiones"]["columnas"].items()}
            if any(v not in (None, "") for v in fila.values()):
                d["revisiones"].append(Revision(**{k: ("" if v is None else v) for k, v in fila.items()}))

        perfil = perfiles.cargar_perfil("EPEN")
        perfil.datos["hojas"]["total"] = 7          # el ETAP dice "1 de 7"
        with tempfile.TemporaryDirectory() as t:
            ruta = excel_writer.escribir_xlsx(Proyecto("EPEN", caratula=Caratula(**d)), perfil, Path(t) / "c.xlsx")
            nuevo = openpyxl.load_workbook(ruta)["Caratula"]
        celdas = set()
        for ws in (orig, nuevo):
            for fila in ws["A1:AD50"]:
                celdas.update(x.coordinate for x in fila if x.value not in (None, ""))
        difs = [(x, orig[x].value, nuevo[x].value) for x in sorted(celdas)
                if txt(orig[x].value).strip() != txt(nuevo[x].value).strip()]
        self.assertEqual(difs, [], f"{len(difs)} celdas distintas")
        self.assertEqual(sorted(str(m) for m in orig.merged_cells.ranges),
                         sorted(str(m) for m in nuevo.merged_cells.ranges))


if __name__ == "__main__":
    unittest.main()
