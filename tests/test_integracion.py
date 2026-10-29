"""La app con las tres partes: carátula + listado A-G + SCADA H-O, y su exportación a un solo Excel."""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

try:
    import tkinter as tk
except ImportError:  # pragma: no cover
    tk = None

from modelo import Caratula, Campo, Funcion, Nivel, Proyecto, Senal  # noqa: E402


class Falso:
    """Reemplaza un diálogo modal: no abre ventana y devuelve el resultado indicado."""
    resultado = {"DI": 200, "AI": 200, "DO": 300}

    def __init__(self, *a, **k):
        pass


@unittest.skipIf(tk is None, "sin tkinter")
class TestApp(unittest.TestCase):
    def setUp(self):
        from app import principal
        self.principal = principal
        try:
            self.app = principal.App()
        except tk.TclError:
            self.skipTest("sin pantalla")
        self.app.withdraw()
        self.mensajes: list[tuple] = []
        from tkinter import filedialog, messagebox
        self._guardado = (filedialog.asksaveasfilename, messagebox.showinfo, messagebox.showerror, messagebox.showwarning,
                          principal.DialogoOffsets, principal.mostrar_texto)
        messagebox.showinfo = lambda *a, **k: self.mensajes.append(a)
        messagebox.showerror = lambda *a, **k: self.mensajes.append(a)
        messagebox.showwarning = lambda *a, **k: self.mensajes.append(a)
        principal.DialogoOffsets = Falso
        principal.mostrar_texto = lambda *a, **k: self.mensajes.append(("resumen", a[2]))
        self.app.wait_window = lambda w: None   # los diálogos falsos no son ventanas reales
        self.addCleanup(self._restaurar)
        self.addCleanup(self.app.destroy)

    def _restaurar(self):
        from tkinter import filedialog, messagebox
        (filedialog.asksaveasfilename, messagebox.showinfo, messagebox.showerror, messagebox.showwarning,
         self.principal.DialogoOffsets, self.principal.mostrar_texto) = self._guardado

    def _campo(self):
        c = Campo(numero="04", descripcion="Campo 04 33KV", prefijo="204LC", controlador="SEL-2241-33kV", senales=[
            Senal(funcion=Funcion.ALARMAS, tag_hmi="A", desc_hmi="204LC Seccionador 229SB04 Falta Tension Auxiliar"),
            Senal(funcion=Funcion.ALARMAS, tag_hmi="B", desc_hmi="204LC Campo 04 33KV QC4 ACTUADA"),
            Senal(funcion=Funcion.SENALIZACION, tag_hmi="C", desc_hmi="204LC Interruptor 25204 Abierto/Cerrado"),
            Senal(funcion=Funcion.SENALIZACION, segunda_fila=True),
            Senal(funcion=Funcion.COMANDO, tag_hmi="D", desc_hmi="204LC Interruptor 25204 Apertura/Cierre"),
            Senal(funcion=Funcion.RESERVA, tag_hmi="N/A", desc_hmi="RESERVA Campo 04", es_reserva=True)])
        self.app.niveles[:] = [Nivel("33KV", "2", [c])]
        self.app.refrescar()
        return c

    # ------------------------------------------------------------ pestañas y columnas
    def test_tres_partes_en_la_ventana(self):
        textos = [self.app.pestanas.tab(i, "text") for i in range(self.app.pestanas.index("end"))]
        self.assertEqual(textos, ["Listado", "Carátula"])
        self.assertEqual(list(self.app.arbol["columns"]), ["item", "nivel", "descrip", "controlador", "funcion", "tag_hmi", "desc_hmi"])

    def test_generar_scada_muestra_h_a_o(self):
        c = self._campo()
        self.app.generar_scada()
        self.assertEqual(len(self.app.arbol["columns"]), 15)
        self.assertEqual(self.app.offsets_iniciales, {"DI": 200, "AI": 200, "DO": 300})
        f = c.senales
        self.assertEqual((f[0].tag_scada, f[0].offset, f[0].s1), ("204.101", 200, "ACTUO"))
        self.assertEqual((f[2].tag_scada, f[2].offset, f[3].offset, f[4].tag_scada, f[4].offset),
                         ("204.122C", 202, 203, "204.122C", 202))          # doble bit en 2 filas; el comando comparte
        self.assertEqual(f[5].tag_scada, "N/A")
        valores = [self.app.arbol.item(i, "values") for i in self.app.arbol.get_children()]
        self.assertEqual(valores[0][7], "204.101")
        self.assertIn("resumen", [m[0] for m in self.mensajes])
        self.assertTrue(any("propuesta" in str(self.app.arbol.item(i, "tags")) for i in self.app.arbol.get_children()))

    def test_cancelar_los_offsets_no_genera(self):
        self._campo()
        self.principal.DialogoOffsets = type("Cancela", (Falso,), {"resultado": None})
        self.app.generar_scada()
        self.assertEqual(len(self.app.arbol["columns"]), 7)

    def test_sin_filas_avisa(self):
        self.app.generar_scada()
        self.assertEqual(self.mensajes[0][0], "Sin filas")

    # ------------------------------------------------------------ edición y propuestas
    def test_lo_editado_en_scada_queda_fijo(self):
        c = self._campo()
        self.app.generar_scada()
        nivel, s = self.app.niveles[0], c.senales[0]
        self.app._asignar(nivel, c, s, "desc_scada", "204LC MI DESC")
        self.app._asignar(nivel, c, s, "offset", "500")
        self.assertEqual((s.desc_scada, s.desc_scada_origen, s.offset), ("204LC MI DESC", "usuario", 500))
        self.assertIn("desc_scada", s.fijos)
        with self.assertRaises(ValueError):
            self.app._asignar(nivel, c, s, "offset", "abc")
        with self.assertRaises(ValueError):
            self.app._asignar(nivel, c, s, "desc_scada", "x" * 21)
        self.app.offsets_iniciales = {"DI": 900, "AI": 200}
        self.app.generar_scada()                         # volver a generar no pisa lo editado a mano
        self.assertEqual((s.desc_scada, s.offset), ("204LC MI DESC", 500))

    def test_confirmar_propuestas_se_recuerdan(self):
        from app import scada_ui
        from scada import Glosario
        c = self._campo()
        self.app.generar_scada()
        pend = scada_ui.pendientes_de_confirmar(self.app.niveles)
        self.assertTrue(pend)
        s = pend[0]
        g = Glosario()
        todas = [x for n in self.app.niveles for cc in n.campos for x in cc.senales]
        n = scada_ui.confirmar_propuestas([(s, "204LC MI ABREV")], todas, g, 20)
        self.assertEqual((n, s.desc_scada, s.desc_scada_origen), (1, "204LC MI ABREV", "usuario"))
        self.assertEqual(list(g.usuario["texto"].values()), ["MI ABREV"])
        self.assertNotIn(s, scada_ui.pendientes_de_confirmar(self.app.niveles))
        with self.assertRaises(ValueError):
            scada_ui.confirmar_propuestas([(s, "x" * 21)], todas, g, 20)
        with self.assertRaises(ValueError):
            scada_ui.confirmar_propuestas([(s, "  ")], todas, g, 20)

    def test_confirmar_propaga_al_comando(self):
        from app import scada_ui
        from scada import Glosario
        c = self._campo()
        self.app.generar_scada()
        sen, cmd = c.senales[2], c.senales[4]
        self.assertEqual(sen.desc_scada, cmd.desc_scada)
        todas = [x for n in self.app.niveles for cc in n.campos for x in cc.senales]
        scada_ui.confirmar_propuestas([(sen, "204LC INT NUEVO")], todas, Glosario(), 20)
        self.assertEqual(cmd.desc_scada, "204LC INT NUEVO")

    # ------------------------------------------------------------ guardar / abrir / exportar
    def test_proyecto_incluye_caratula_y_offsets(self):
        self._campo()
        self.app.caratula_ui.vars["obra"].set("LAT de prueba")
        self.app.offsets_iniciales = {"DI": 10}
        p = self.app._proyecto()
        self.assertEqual((p.caratula.obra, p.offsets_iniciales), ("LAT de prueba", {"DI": 10}))
        with tempfile.TemporaryDirectory() as t:
            ruta = Path(t) / "p.json"
            p.guardar(ruta)
            q = Proyecto.cargar(ruta)
            self.assertEqual(q.caratula.obra, "LAT de prueba")

    def test_exporta_las_tres_partes(self):
        from openpyxl import load_workbook
        from tkinter import filedialog
        c = self._campo()
        self.app.generar_scada()
        self.app.caratula_ui.vars["obra"].set("LAT de prueba")
        self.app.caratula_ui.vars["cliente"].set("EPEN")
        with tempfile.TemporaryDirectory() as t:
            ruta = Path(t) / "x.xlsx"
            filedialog.asksaveasfilename = lambda **k: str(ruta)
            self.app.exportar()
            self.assertEqual(self.mensajes[-1][0], "Exportado")
            self.assertIn("carátula, listado A-G, SCADA H-O", self.mensajes[-1][1])
            wb = load_workbook(ruta)
            self.assertEqual(wb.sheetnames, ["Caratula", "EPEN"])
            self.assertEqual(wb["Caratula"]["C34"].value, "LAT de prueba")
            ws = wb["EPEN"]
            self.assertEqual(ws.max_column, 15)
            self.assertEqual((ws["H6"].value, ws["I6"].value, ws["K6"].value, ws["L6"].value, ws["O6"].value),
                             ("204.101", 200, "NORMAL", "ACTUO", "N/A"))
            self.assertEqual((ws["H8"].value, ws["I8"].value, ws["I9"].value), ("204.122C", 202, 203))

    def test_exporta_solo_listado_sin_caratula_ni_scada(self):
        from openpyxl import load_workbook
        from tkinter import filedialog
        self._campo()
        with tempfile.TemporaryDirectory() as t:
            ruta = Path(t) / "x.xlsx"
            filedialog.asksaveasfilename = lambda **k: str(ruta)
            self.app.exportar()
            self.assertEqual(self.mensajes[-1][0], "Exportado")
            wb = load_workbook(ruta)
            self.assertEqual((wb.sheetnames, wb["EPEN"].max_column), (["EPEN"], 7))

    def test_caratula_invalida_impide_exportar(self):
        from tkinter import filedialog
        self._campo()
        self.app.caratula_ui.vars["obra"].set("X")
        self.app.caratula_ui.vars["numero_documento"].set("A" * 25)
        with tempfile.TemporaryDirectory() as t:
            filedialog.asksaveasfilename = lambda **k: str(Path(t) / "x.xlsx")
            self.app.exportar()
            self.assertEqual(self.mensajes[-1][0], "Revisá la carátula")
            self.assertFalse((Path(t) / "x.xlsx").exists())


if __name__ == "__main__":
    unittest.main()
