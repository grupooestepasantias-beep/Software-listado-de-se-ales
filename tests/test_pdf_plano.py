import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import excel_writer
import pdf_plano
import perfiles
from modelo import Funcion, Proyecto

PLANO = Path(r"\\192.168.15.120\PasantiasMM\Sofware listado\Ejemplos\ETAP-0325-ET-CT-PL-EF-6402_1.pdf")

try:
    import pymupdf  # noqa: F401
except ImportError:
    pymupdf = None


class TestClasificacion(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.perfil = perfiles.cargar_perfil("EPEN")
        cls.cfg = pdf_plano.cargar_config(cls.perfil)

    def banda(self, textos, pagina=1):
        return pdf_plano._senales_de_banda(textos, self.cfg, pagina)

    def test_ignora_alimentacion_y_marca_reservas(self):
        r = self.banda(["Negativo 48 VCC", "Reserva", "Positivo 110 Vcc", "Reserva"])
        self.assertEqual([s.funcion for s in r], [Funcion.RESERVA, Funcion.RESERVA])

    def test_pares_de_doble_bit_en_orden_de_estados(self):
        r = self.banda(["Seccionador 229SL04 Cerrado", "Seccionador 229SL04 Abierto",
                        "Seccionador 229SL04 Remoto", "Seccionador 229SL04 Local"])
        self.assertEqual([s.funcion for s in r], [Funcion.SENALIZACION, Funcion.SENALIZACION])
        # la 1ª fila es el 1er estado (Abierto / Local), aunque en el plano figure después
        self.assertEqual(r[0].columnas, ["Seccionador 229SL04 Abierto", "Seccionador 229SL04 Cerrado"])
        self.assertEqual(r[0].texto, "Seccionador 229SL04 Abierto/Cerrado")
        self.assertEqual(r[1].columnas[0], "Seccionador 229SL04 Local")

    def test_no_empareja_equipos_distintos(self):
        r = self.banda(["Seccionador 229SL04 Abierto", "Seccionador 229SB04 Cerrado"])
        self.assertEqual(len(r), 2)
        self.assertTrue(all(s.funcion == Funcion.ALARMAS for s in r))

    def test_banda_de_salidas_son_comandos(self):
        r = self.banda(["Positivo 110 Vcc", "Seccionador 229SL04 Cierre", "Positivo 110 Vcc",
                        "Seccionador 229SL04 Apertura", "Reserva", "Interruptor 25204 Cierre"])
        self.assertEqual([s.funcion for s in r], [Funcion.COMANDO, Funcion.RESERVA, Funcion.COMANDO])
        self.assertEqual(r[0].columnas[0], "Seccionador 229SL04 Apertura")

    def test_codigo_de_nivel_por_tension(self):
        esperado = {"132KV": "3", "66KV": "3", "33KV": "2", "13,2KV": "1", "13KV": "1", "6.6KV": "1",
                    "6,6 kV": "1", "SSAA": ""}
        for nombre, cod in esperado.items():
            self.assertEqual(pdf_plano.codigo_nivel(nombre, self.cfg), cod, nombre)

    def test_controlador_por_defecto(self):
        self.assertEqual(pdf_plano.controlador_por_defecto("33KV", self.cfg), "SEL-2241-33kV")
        self.assertEqual(pdf_plano.controlador_por_defecto("132 KV", self.cfg), "SEL-2241-132kV")
        self.assertEqual(pdf_plano.controlador_por_defecto("6.6KV", self.cfg), "SEL-2241-6,6kV")  # coma, como el ETAP

    def test_la_tension_del_controlador_sigue_al_nivel(self):
        f = pdf_plano.con_tension_del_nivel
        self.assertEqual(f("SEL-2241-33kV", "132KV"), "SEL-2241-132kV")
        self.assertEqual(f("SEL-2241-132kV", "6,6KV"), "SEL-2241-6,6kV")
        self.assertEqual(f("SCHNEIDER P142_C04", "132KV"), "SCHNEIDER P142_C04")  # un relé no lleva tensión
        self.assertEqual(f("SEL-3555", "132KV"), "SEL-3555")
        self.assertEqual(f("SEL-2241-33kV", "RESERVAS"), "SEL-2241-33kV")  # nivel sin tensión: no cambia

    def test_las_senales_de_proteccion_llevan_el_rele_del_campo(self):
        niv = self.armar(["Prot. Campo 04 33 KV Falla Interna", "Interruptor 25204 Falta Fuerza Motriz",
                          "Campo 04 33KV Falla Multimedidor", "2TV04- Falta Vca Prot"],
                         destinos={5: ("33KV", "04", "", "", "AP", "SCHNEIDER P142_C04")})
        por = {s.desc_hmi: s.controlador for s in niv[0].campos[0].senales}
        self.assertEqual(por["204AP Proteccion Campo 04 33 KV Falla Interna"], "SCHNEIDER P142_C04")
        self.assertEqual(por["204AP Interruptor 25204 Falta Fuerza Motriz"], "SEL-2241-33kV")
        self.assertEqual(por["204AP Campo 04 33KV Falla Multimedidor"], "SEL-2241-33kV")
        self.assertEqual(por["204AP 2TV04- Falta Vca Proteccion"], "SEL-2241-33kV")  # 'Prot' no está al comienzo

    def armar(self, textos, **kw):
        return pdf_plano.a_niveles([pdf_plano.PaginaLeida(5, "33KV", "04", self.banda(textos, 5))], self.perfil, **kw)

    def test_filas_a_g(self):
        niv = self.armar(["Seccionador 229SL04 Cerrado", "Seccionador 229SL04 Abierto", "Reserva",
                          "Interruptor 25204 Falta Fuerza Motriz"], destinos={5: ("33KV", "04", "Campo 04", "", "AP")})
        c = niv[0].campos[0]
        self.assertEqual((niv[0].nombre, niv[0].codigo, c.descripcion, c.prefijo, c.controlador),
                         ("33KV", "2", "Campo 04", "204AP", "SEL-2241-33kV"))
        a, s1, s2 = c.senales  # alarma primero (orden fijo), luego la señalización con sus 2 filas
        self.assertEqual((a.funcion, a.controlador, a.tag_hmi, a.desc_hmi),
                         (Funcion.ALARMAS, "SEL-2241-33kV", "Interruptor_25204_F_Fza_Motriz",
                          "204AP Interruptor 25204 Falta Fuerza Motriz"))
        self.assertEqual((s1.tag_hmi, s1.desc_hmi), ("Seccionador_229SL04_Abierto_Cerrado",
                                                      "204AP Seccionador 229SL04 Abierto/Cerrado"))
        self.assertEqual((s2.segunda_fila, s2.funcion, s2.controlador, s2.tag_hmi, s2.desc_hmi),
                         (True, Funcion.SENALIZACION, "SEL-2241-33kV", "", ""))

    def test_mediciones_del_campo_al_final_con_el_estilo_del_etap(self):
        senales = (self.banda(["Interruptor 25204 Falta Fuerza Motriz"], 5)
                   + self.banda(["Seccionador 229SL04 Apertura", "Seccionador 229SL04 Cierre"], 5))  # 2 módulos
        pg = pdf_plano.PaginaLeida(5, "33KV", "04", senales, mediciones=True)
        niv = pdf_plano.a_niveles([pg], self.perfil, destinos={5: ("33KV", "04", "", "", "AP")})
        c = niv[0].campos[0]
        self.assertEqual([s.funcion for s in c.senales][:3], [Funcion.ALARMAS, Funcion.COMANDO, Funcion.COMANDO])
        med = c.senales[3:]  # alarma, comando (2 filas) y al final las mediciones
        self.assertEqual(len(med), 13)
        self.assertTrue(all(s.funcion == Funcion.MEDICION and s.controlador == "ION 7400 / SCHNEIDER P142" for s in med))
        self.assertEqual((med[0].desc_hmi, med[0].tag_hmi), ("204AP TENSION CAMPO 04 33kV FASE R",
                                                             "TENSION_CAMPO_04_33kV_FASE_R"))
        self.assertEqual(med[-1].desc_hmi, "204AP CORRIENTE PROMEDIO CAMPO 04 33kV")
        self.assertEqual(med[7].desc_hmi, "204AP ANGULO DE FASE TOTAL CAMPO 04 33kV")

    def test_controlador_con_la_tension_del_nivel(self):
        f = pdf_plano.controlador_con_tension
        self.assertEqual(f("SEL-2241", "33KV"), "SEL-2241-33kV")
        self.assertEqual(f("SEL-2241", "6.6KV"), "SEL-2241-6,6kV")
        self.assertEqual(f("SEL-2241-132kV", "33KV"), "SEL-2241-132kV")  # ya la trae
        self.assertEqual(f("SM_CPU866e", "RESERVAS"), "SM_CPU866e")      # nivel sin tensión

    def test_sin_multimedidor_no_hay_mediciones_salvo_que_se_pidan(self):
        pg = pdf_plano.PaginaLeida(5, "33KV", "04", self.banda(["Interruptor 25204 Falta Fuerza Motriz"], 5))
        self.assertEqual(len(pdf_plano.a_niveles([pg], self.perfil)[0].campos[0].senales), 1)
        pedidas = pdf_plano.a_niveles([pg], self.perfil, destinos={5: ("33KV", "04", "", "", "", "", True)})
        self.assertEqual(len(pedidas[0].campos[0].senales), 14)

    def test_las_mediciones_no_se_repiten_si_el_campo_tiene_varias_paginas(self):
        a = pdf_plano.PaginaLeida(5, "33KV", "04", self.banda(["Eq 1 Falla"], 5), mediciones=True)
        b = pdf_plano.PaginaLeida(6, "33KV", "04", self.banda(["Eq 2 Falla"], 6), mediciones=True)
        c = pdf_plano.a_niveles([a, b], self.perfil)[0].campos[0]
        self.assertEqual(sum(s.funcion == Funcion.MEDICION for s in c.senales), 13)

    def test_reservas_juntas_al_final_con_el_campo_en_la_descripcion(self):
        a = pdf_plano.PaginaLeida(1, "33KV", "04", self.banda(["Reserva", "Eq 1 Falla"], 1))
        b = pdf_plano.PaginaLeida(2, "132KV", "04", self.banda(["Reserva", "Eq 2 Falla"], 2))
        niv = pdf_plano.a_niveles([a, b], self.perfil)
        self.assertEqual([n.nombre for n in niv], ["33KV", "132KV", "RESERVAS"])
        r = niv[-1].campos[0].senales
        self.assertEqual([(s.tag_hmi, s.desc_hmi) for s in r],
                         [("N/A", "RESERVA 33KV Campo 04"), ("N/A", "RESERVA 132KV Campo 04")])


class TestOtrosPlanos(unittest.TestCase):
    """Reglas aprendidas de los otros planos de ingeniería (carpeta ape): hojas de cableado, módulos analógicos."""

    @classmethod
    def setUpClass(cls):
        cls.perfil = perfiles.cargar_perfil("EPEN")
        cls.cfg = pdf_plano.cargar_config(cls.perfil)

    def test_rotulos_de_cable_no_son_senales(self):
        no = ["-X1.6", "H23/4A", "H17/2B", "-P1.1", "E13", "110 Vcc (+)", "FASE R", "1,5mm2/Ne", "2.5 mm2", "-AL1"]
        si = ["Reserva", "RESERVA", "Positivo 110 Vcc", "Negativo 48 VCC", "Interruptor 25204 Cerrado",
              "ALARMA PRESION SF6", "Arrollamiento 132 KV 0-150º C", "Falta tensión auxiliar"]
        for t in no:
            self.assertFalse(pdf_plano._es_rotulo(t, self.cfg), t)
        for t in si:
            self.assertTrue(pdf_plano._es_rotulo(t, self.cfg), t)

    def test_modulo_analogico_son_mediciones_y_sus_reservas_no_se_cuentan(self):
        r = pdf_plano._senales_de_banda(["Arrollamiento 132 KV 0-150º C"] + ["Reserva"] * 15, self.cfg, 55)
        self.assertEqual([(s.funcion, s.texto) for s in r], [(Funcion.MEDICION, "Arrollamiento 132 KV 0-150º C")])
        # con Positivo/Negativo es un módulo digital: ahí las reservas sí son señales
        r = pdf_plano._senales_de_banda(["Negativo 48 VCC", "Reserva", "Falla X 0-1"], self.cfg, 1)
        self.assertEqual([s.funcion for s in r], [Funcion.RESERVA, Funcion.ALARMAS])

    def test_abrir_cerrar_es_un_comando_de_doble_bit(self):
        r = pdf_plano._senales_de_banda(["Interruptor QA Cerrar", "Interruptor QA Abrir"], self.cfg, 1)
        self.assertEqual([(s.funcion, s.texto) for s in r], [(Funcion.COMANDO, "Interruptor QA Abrir/Cerrar")])


@unittest.skipUnless(pymupdf and PLANO.exists(), "falta pymupdf o el plano de ejemplo")
class TestPlanoReal(unittest.TestCase):
    def test_pagina_29_y_exportar_a_g(self):
        perfil = perfiles.cargar_perfil("EPEN")
        paginas = pdf_plano.leer_plano(PLANO, perfil, solo_paginas={29})
        self.assertEqual(len(paginas), 1)
        pg = paginas[0]
        self.assertEqual((pg.nivel, pg.campo), ("33KV", "04"))
        textos = [s.texto for s in pg.senales]
        self.assertIn("Seccionador 229SL04 Abierto/Cerrado", textos)
        self.assertIn("Interruptor 25204 Apertura/Cierre", textos)
        self.assertFalse(any("Positivo" in t or "Negativo" in t for t in textos))
        self.assertEqual(sum(s.funcion == Funcion.COMANDO for s in pg.senales), 3)
        self.assertTrue(pg.mediciones)  # la página nombra a un multimedidor ("Falla Multimedidor")

        niv = pdf_plano.a_niveles(paginas, perfil)
        self.assertEqual(sum(s.funcion == Funcion.MEDICION for s in niv[0].campos[0].senales), 13)
        p = Proyecto(perfil_id="EPEN", niveles=niv, item_inicial=1)
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("falta openpyxl")
        with tempfile.TemporaryDirectory() as t:
            ruta = excel_writer.escribir_xlsx(p, perfil, Path(t) / "s.xlsx", incluir_caratula=False,
                                              hasta_columna="G")
            wb = load_workbook(ruta)
            self.assertEqual(wb.sheetnames, ["EPEN"])
            ws = wb["EPEN"]
            n = len(p.todas_las_senales())
            self.assertEqual(ws.cell(5 + n, 1).value, n)  # ITEM correlativo hasta la última fila
            self.assertEqual((ws["B6"].value, ws["D6"].value, ws["E6"].value),
                             ("33KV", "SEL-2241-33kV", "ALARMAS"))
            self.assertEqual(ws.max_column, 7)  # solo A-G
            self.assertEqual(ws.auto_filter.ref, f"A5:G{5 + n}")
            self.assertEqual(ws.print_area, f"'EPEN'!$A$1:$G${5 + n}")


APE = Path(r"\\192.168.15.120\PasantiasMM\ape")
TC6302 = APE / "TC3-04 TRAFO TR02" / "ETAP-0325-ET-CT-PL-EF-6302_1_CAF.pdf"
TP6311 = APE / "TP3-06 TRAFO TR03 (Puerta vidrio)" / "como esta cableado" / "ETAP-0325-ET-CT-PL-EF-6311_B.pdf"
PARA_CAO = APE / "TC3-04 TRAFO TR02" / "ETAP-0325-ET-CT-PL-EF-6302 para CAO.pdf"


@unittest.skipUnless(pymupdf and TC6302.exists(), "falta pymupdf o la carpeta ape")
class TestPlanosDeLaCarpetaApe(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.perfil = perfiles.cargar_perfil("EPEN")

    def test_tablero_de_control_con_modulo_analogico_y_cajetin(self):
        pg = {p.pagina: p for p in pdf_plano.leer_plano(TC6302, self.perfil, solo_paginas={53, 54, 55})}
        self.assertEqual(sorted(pg), [53, 54, 55])
        for p in pg.values():  # el nivel y el campo salen del cajetín (también en las hojas sin "Campo NN" en los rótulos)
            self.assertEqual((p.nivel, p.campo), ("132KV", "04"))
        self.assertEqual(pg[55].descripcion, "TRANSF TR02 132Kv")
        self.assertEqual([(s.funcion, s.texto) for s in pg[55].senales],
                         [(Funcion.MEDICION, "Arrollamiento 132 KV 0-150º C")])
        self.assertTrue(any(s.texto.startswith("Prot. principal trafo") for s in pg[54].senales))

    def test_hojas_de_cableado_de_un_tablero_de_proteccion_no_dan_senales(self):
        # entradas binarias, alimentación, disparos…: son cableado del relé, no señales hacia el control
        self.assertEqual(pdf_plano.leer_plano(TP6311, self.perfil, solo_paginas={24, 27, 2, 6, 12, 33, 35}), [])

    def test_tiras_de_senalizacion_del_rele_al_control(self):
        pg = {p.pagina: p for p in pdf_plano.leer_plano(TP6311, self.perfil, solo_paginas={29, 30, 31})}
        self.assertEqual(sorted(pg), [29, 30, 31])
        for p in pg.values():
            self.assertEqual((p.origen, p.nivel, p.campo), ("rele", "132KV", "06"))  # TP3 06 -> 132KV, campo 06
        self.assertEqual([s.texto for s in pg[29].senales],
                         ["FALTA ALIMENTAC. PROT. PRINCIPAL", "FALTA ALIMENTAC. PROT. RESPALDO",
                          "FALTA ALIMENTAC. PROT. R.NEUTRO", "BAJA PRESIÓN BLOQUEO SF6",
                          "FALLA CARGA RESORTE", "BLOQUEO AL CIERRE"])  # 'ALARMAS' es una cabecera, no una señal
        textos31 = [s.texto for s in pg[31].senales]
        self.assertIn("PROTECCION PRINCIPAL DISPARO GENERAL", textos31)  # la cabecera del grupo se antepone
        self.assertIn("REACTOR DE NEUTRO BUCHHOLZ", textos31)
        self.assertTrue(all(s.de_rele and s.funcion == Funcion.ALARMAS for p in pg.values() for s in p.senales))

    def test_las_senales_del_rele_llevan_el_rele_del_campo_con_descripcion_completa(self):
        pg = pdf_plano.leer_plano(TP6311, self.perfil, solo_paginas={29})
        niv = pdf_plano.a_niveles(pg, self.perfil, {29: ("132KV", "06", "", "SEL-2241-132kV", "AP", "SCHNEIDER P543_C06")})
        filas = niv[0].campos[0].senales
        self.assertEqual({s.controlador for s in filas}, {"SCHNEIDER P543_C06"})
        self.assertEqual(filas[0].desc_hmi, "306AP FALTA Alimentacion PROT. PRINCIPAL".replace("PROT.", "Proteccion"))
        # sin relé cargado se usa el controlador del campo
        sin = pdf_plano.a_niveles(pg, self.perfil, {29: ("132KV", "06", "", "SEL-2241-132kV", "AP", "")})
        self.assertEqual({s.controlador for s in sin[0].campos[0].senales}, {"SEL-2241-132kV"})

    def test_la_impresion_de_un_listado_de_senales_no_es_un_esquema(self):
        listado = APE / "En ejecucion" / "ETAP-0325-ET-CT-MC-LI-6209_ LISTADO DE SEÑALES.pdf"
        if not listado.exists():
            self.skipTest("falta el listado en PDF")
        with self.assertRaises(pdf_plano.PdfNoEsquema):
            pdf_plano.leer_plano(listado, self.perfil)

    def test_los_pdf_en_ejecucion_se_leen(self):
        for nombre, campo in (("ETAP-0325-ET-CT-PL-EF-6402_D.pdf", "04"), ("ETAP-0325-ET-CT-PL-EF-6303_C.pdf", "06")):
            ruta = APE / "En ejecucion" / nombre
            if not ruta.exists():
                self.skipTest(f"falta {nombre}")
            pg = pdf_plano.leer_plano(ruta, self.perfil)
            self.assertTrue(pg and any(s.funcion == Funcion.ALARMAS for s in pg[0].senales), nombre)
            self.assertEqual(pg[0].campo, campo)

    @unittest.skipUnless(PARA_CAO.exists(), "falta el PDF 'para CAO'")
    def test_pdf_sin_texto_avisa_en_vez_de_devolver_vacio(self):
        with self.assertRaises(pdf_plano.PdfSinTexto):
            pdf_plano.leer_plano(PARA_CAO, self.perfil)


if __name__ == "__main__":
    unittest.main()
