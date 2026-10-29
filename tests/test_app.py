import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import excel_writer
import perfiles
from app import validacion as v

try:
    import tkinter as tk
except ImportError:
    tk = None


def _datos_ok() -> dict:
    d = v.datos_vacios()
    d.update(cliente="EPEN", obra="LAT 132 kV", titulo="ETAP - LISTADO DE SEÑALES",
             numero_documento="ETAP-0722-ET-CT-MC-LI-6209", item_inicial="1",
             offsets={"DI": "200", "AI": "200", "DO": "0", "AO": "0"},
             referencias=[("ETAP-0722-ET-CT-PL-EF-6106", "Esquema"), ("", "")],
             revisiones=[{"rev": "0", "fecha": "21/04/2025", "descripcion": "APC", "proyecto": "DV",
                          "revisor": "GG", "aprobador": "ME"}])
    return d


class TestValidacion(unittest.TestCase):
    def test_formulario_valido(self):
        self.assertEqual(v.validar(_datos_ok(), 14, 4, 20), [])

    def test_nada_es_obligatorio_por_ahora(self):
        d = v.datos_vacios()  # ITEM arranca en 1; la carátula y los OFFSET (SCADA) se completan después
        self.assertEqual(v.validar(d, 14, 4, 20), [])
        d["item_inicial"] = ""  # vacío tampoco bloquea: se toma 1
        self.assertEqual(v.validar(d, 14, 4, 20), [])
        self.assertEqual(v.a_proyecto(d, "EPEN").item_inicial, 1)

    def test_item_no_numerico_o_negativo(self):
        d = _datos_ok()
        d["item_inicial"] = "-3"
        self.assertTrue(any("ITEM" in m for m in v.validar(d, 14, 4, 20)))
        d["item_inicial"] = "abc"
        self.assertTrue(any("ITEM" in m for m in v.validar(d, 14, 4, 20)))

    def test_largo_documento(self):
        d = _datos_ok()
        d["numero_documento"] = "ETAP-0722"  # más corto: permitido
        self.assertEqual(v.validar(d, 14, 4, 20), [])
        d["numero_documento"] = "ETAP-0722-ET-CT-MC-LI-62099"  # 21 caracteres sin guiones
        self.assertTrue(any("Número de documento" in m for m in v.validar(d, 14, 4, 20)))

    def test_limites_y_filas_vacias(self):
        d = _datos_ok()
        d["referencias"] = [(f"P{i}", "x") for i in range(15)]
        self.assertTrue(any("referencias" in m for m in v.validar(d, 14, 4, 20)))
        d["referencias"] = [("", ""), ("", "")]  # filas vacías no cuentan
        self.assertEqual(v.validar(d, 14, 4, 20), [])

    def test_ida_y_vuelta_proyecto(self):
        p = v.a_proyecto(_datos_ok(), "EPEN")
        self.assertEqual(p.item_inicial, 1)
        self.assertEqual(p.offsets_iniciales["DI"], 200)
        self.assertEqual(len(p.caratula.referencias), 1)
        d2 = v.desde_proyecto(p)
        self.assertEqual(d2["numero_documento"], "ETAP-0722-ET-CT-MC-LI-6209")
        self.assertEqual(d2["revisiones"][0]["rev"], "0")
        with tempfile.TemporaryDirectory() as t:
            p.guardar(Path(t) / "p.json")
            from modelo import Proyecto
            self.assertEqual(Proyecto.cargar(Path(t) / "p.json").to_dict(), p.to_dict())

    def test_exporta_excel_desde_formulario(self):
        perfil = perfiles.cargar_perfil("EPEN")
        p = v.a_proyecto(_datos_ok(), "EPEN")
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("falta openpyxl")
        with tempfile.TemporaryDirectory() as t:
            ruta = excel_writer.escribir_xlsx(p, perfil, Path(t) / "x.xlsx")
            ws = load_workbook(ruta)["Caratula"]
            self.assertEqual(ws["C34"].value, "LAT 132 kV")
            self.assertEqual(ws["D44"].value, "E")


@unittest.skipIf(tk is None, "sin tkinter")
class TestVentana(unittest.TestCase):
    """La ventana nueva: tabla A-G, edición de celdas, mover/borrar bloques y exportar solo A-G."""

    def setUp(self):
        from app.principal import App
        try:
            self.app = App()
        except tk.TclError:
            self.skipTest("sin pantalla")
        self.app.withdraw()

    def tearDown(self):
        self.app.destroy()

    def _campo(self):
        from modelo import Campo, Funcion, Nivel, Senal
        c = Campo(numero="04", descripcion="Campo 04", controlador="SEL-2241-33kV", senales=[
            Senal(funcion=Funcion.ALARMAS, tag_hmi="A", desc_hmi="204 A", controlador="SEL-2241-33kV"),
            Senal(funcion=Funcion.SENALIZACION, tag_hmi="B", desc_hmi="204 B"),
            Senal(funcion=Funcion.SENALIZACION, segunda_fila=True),
            Senal(funcion=Funcion.COMANDO, tag_hmi="C")])
        self.app.niveles[:] = [Nivel("33KV", "2", [c])]
        self.app.refrescar()
        return c

    def test_siete_columnas_y_celdas_combinadas(self):
        self._campo()
        self.assertEqual(list(self.app.arbol["columns"]), ["item", "nivel", "descrip", "controlador", "funcion", "tag_hmi", "desc_hmi"])
        filas = [self.app.arbol.item(i, "values") for i in self.app.arbol.get_children()]
        self.assertEqual([str(f[0]) for f in filas], ["1", "2", "3", "4"])
        self.assertEqual((filas[0][1], filas[0][2], filas[1][1], filas[1][2]), ("33KV", "Campo 04", "", ""))

    def test_item_inicial_renumera(self):
        self._campo()
        self.app.item.set("100")
        filas = [str(self.app.arbol.item(i, "values")[0]) for i in self.app.arbol.get_children()]
        self.assertEqual(filas, ["100", "101", "102", "103"])

    def test_editar_celdas_y_grupos(self):
        from modelo import Funcion
        c = self._campo()
        nivel = self.app.niveles[0]
        self.app._asignar(nivel, c, c.senales[0], "nivel", "132KV")
        self.assertEqual((nivel.nombre, nivel.codigo), ("132KV", "3"))
        self.app._asignar(nivel, c, c.senales[0], "descrip", "LAT Loma Campana")
        self.app._asignar(nivel, c, c.senales[0], "funcion", "reserva")
        self.assertEqual((c.descripcion, c.senales[0].funcion, c.senales[0].es_reserva),
                         ("LAT Loma Campana", Funcion.RESERVA, True))
        with self.assertRaises(ValueError):
            self.app._asignar(nivel, c, c.senales[0], "funcion", "otra cosa")

    def test_cambiar_el_nivel_cambia_la_tension_del_controlador(self):
        c = self._campo()
        nivel = self.app.niveles[0]
        c.senales[0].controlador = "SCHNEIDER P142_C04"  # un relé no lleva tensión
        self.app._asignar(nivel, c, c.senales[0], "nivel", "132KV")
        self.assertEqual((c.controlador, c.senales[1].controlador), ("SEL-2241-132kV", ""))
        self.assertEqual(c.senales[0].controlador, "SCHNEIDER P142_C04")
        self.app._asignar(nivel, c, c.senales[0], "nivel", "6,6KV")
        self.assertEqual(c.controlador, "SEL-2241-6,6kV")

    def test_agregar_mediciones_al_final_del_campo_sin_repetir(self):
        from tkinter import messagebox
        from modelo import Funcion
        c = self._campo()
        avisos = []
        messagebox.showinfo = lambda *a, **k: avisos.append(a)
        self.app.arbol.selection_set(list(self.app.arbol.get_children())[0])
        self.app.agregar_mediciones()
        med = [s for s in c.senales if s.funcion == Funcion.MEDICION]
        self.assertEqual((len(med), c.senales[-1].funcion), (13, Funcion.MEDICION))
        self.assertEqual((med[0].desc_hmi, med[0].controlador), ("204 TENSION CAMPO 04 33kV FASE R",
                                                                 "ION 7400 / SCHNEIDER P142"))
        self.app.arbol.selection_set(list(self.app.arbol.get_children())[0])
        self.app.agregar_mediciones()  # ya las tiene: no se duplican
        self.assertEqual(sum(s.funcion == Funcion.MEDICION for s in c.senales), 13)
        self.assertEqual(len(avisos), 1)

    def _con_dialogo_controlador(self, resultado):
        """Reemplaza la ventana modal del controlador por una que responde `resultado`."""
        from app import principal
        recibido = {}

        class Falso:
            def __init__(self, master, actual, nivel, entradas):
                recibido.update(actual=actual, nivel=nivel, modelos=len(entradas))
                self.resultado = resultado

        self.addCleanup(setattr, principal, "DialogoControlador", principal.DialogoControlador)  # no se filtra
        principal.DialogoControlador = Falso
        self.app.wait_window = lambda w: None
        return recibido

    def test_el_controlador_se_elige_del_catalogo_para_la_fila_o_el_campo(self):
        c = self._campo()
        ids = list(self.app.arbol.get_children())
        recibido = self._con_dialogo_controlador(("SEL-2241-33kV", "fila"))
        self.app._elegir_controlador(self.app.niveles[0], c, c.senales[1])
        self.assertEqual([s.controlador for s in c.senales], ["SEL-2241-33kV", "SEL-2241-33kV", "", ""])
        self.assertEqual((recibido["nivel"], recibido["modelos"] >= 40), ("33KV", True))
        self._con_dialogo_controlador(("SM_CPU866e", "campo"))
        self.app._elegir_controlador(self.app.niveles[0], c, c.senales[0])
        self.assertEqual({s.controlador for s in c.senales}, {"SM_CPU866e"})
        self._con_dialogo_controlador(None)  # cancelar no cambia nada
        self.app._elegir_controlador(self.app.niveles[0], c, c.senales[0])
        self.assertEqual({s.controlador for s in c.senales}, {"SM_CPU866e"})

    def test_la_ventana_del_controlador_propone_la_tension_solo_para_controladores(self):
        from app.principal import DialogoControlador
        import catalogo
        v = DialogoControlador(self.app, "", "33KV", catalogo.cargar()["entradas"])
        try:
            self.assertIn("SEL-2241", v.combo["values"])
            self.assertIn("SM_DI32", v.combo["values"])
            v.valor.set("SEL-2241")
            v._elegido()
            self.assertTrue(v.con_tension.get())  # controlador: sugiere SEL-2241-33kV
            v.alcance.set("campo")
            v._ok()
            self.assertEqual(v.resultado, ("SEL-2241-33kV", "campo"))
        finally:
            v.destroy()
        v = DialogoControlador(self.app, "", "33KV", catalogo.cargar()["entradas"])
        try:
            v.valor.set("SEL-2244-2")  # módulo de E/S: sin tensión
            v._elegido()
            self.assertFalse(v.con_tension.get())
            v._ok()
            self.assertEqual(v.resultado, ("SEL-2244-2", "fila"))
        finally:
            v.destroy()
        v = DialogoControlador(self.app, "", "RESERVAS", [])  # nivel sin tensión: la casilla no aplica
        try:
            v.valor.set("ACME 9000")  # se puede escribir un modelo que no está en el catálogo
            v._ok()
            self.assertEqual(v.resultado, ("ACME 9000", "fila"))
        finally:
            v.destroy()

    def test_usar_un_modelo_del_catalogo(self):
        c = self._campo()
        ids = list(self.app.arbol.get_children())
        self.app.arbol.selection_set(ids[1])
        self.assertEqual(self.app._aplicar_controlador("SEL-2241", True, "filas"), 1)  # con la tension del nivel
        self.assertEqual([s.controlador for s in c.senales], ["SEL-2241-33kV", "SEL-2241-33kV", "", ""])
        self.app.arbol.selection_set(list(self.app.arbol.get_children())[0])
        self.assertEqual(self.app._aplicar_controlador("SM_CPU866e", False, "campo"), 4)  # todo el campo, sin tension
        self.assertEqual({s.controlador for s in c.senales}, {"SM_CPU866e"})
        self.assertEqual(self.app._aplicar_controlador("SEL-2241", True, "filas"), 1)  # se conserva la fila elegida
        self.assertEqual(c.senales[0].controlador, "SEL-2241-33kV")
        self.app.arbol.selection_remove(*self.app.arbol.selection())
        self.assertEqual(self.app._aplicar_controlador("SEL-2241", True, "filas"), 0)  # sin seleccion: nada

    def test_la_ventana_del_catalogo_lista_filtra_y_aplica(self):
        from app.catalogo_ui import VentanaCatalogo
        llamadas = []
        v = VentanaCatalogo(self.app, lambda m, t, a: llamadas.append((m, t, a)) or 1)
        try:
            self.assertGreater(len(v.arbol.get_children()), 40)
            v.texto.set("2244-3")
            self.assertEqual([v.arbol.item(i, "values")[3] for i in v.arbol.get_children()], ["SEL-2244-3"])
            v.arbol.selection_set(v.arbol.get_children()[0])
            v.alcance.set("campo")
            v._usar()
            self.assertEqual(llamadas, [("SEL-2244-3", True, "campo")])
            v.marca.set("Schneider Electric")
            v.texto.set("")
            self.assertTrue(all(v.arbol.item(i, "values")[0] == "Schneider Electric" for i in v.arbol.get_children()))
        finally:
            v.destroy()

    def test_doble_bit_se_mueve_y_se_borra_como_bloque(self):
        from tkinter import messagebox
        c = self._campo()
        ids = list(self.app.arbol.get_children())
        self.app.arbol.selection_set(ids[1])
        self.app.mover(1)
        self.assertEqual([s.tag_hmi for s in c.senales], ["A", "C", "B", ""])
        messagebox.askyesno = lambda *a, **k: True
        self.app.arbol.selection_set(list(self.app.arbol.get_children())[2])  # la señal B
        self.app.borrar()
        self.assertEqual([s.tag_hmi for s in c.senales], ["A", "C"])

    def test_agregar_fila_despues_del_bloque(self):
        c = self._campo()
        self.app.arbol.selection_set(list(self.app.arbol.get_children())[1])
        self.app.agregar_fila()
        self.assertEqual(len(c.senales), 5)
        self.assertEqual((c.senales[3].controlador, c.senales[3].segunda_fila), ("SEL-2241-33kV", False))

    def test_exportar_solo_a_g(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("falta openpyxl")
        from tkinter import filedialog, messagebox
        self._campo()
        with tempfile.TemporaryDirectory() as t:
            ruta = Path(t) / "x.xlsx"
            filedialog.asksaveasfilename = lambda **k: str(ruta)
            mensajes = []
            messagebox.showinfo = lambda *a, **k: mensajes.append(a)
            messagebox.showerror = lambda *a, **k: mensajes.append(a)
            self.app.exportar()
            self.assertEqual([m[0] for m in mensajes], ["Exportado"])
            ws = load_workbook(ruta)["EPEN"]
            self.assertEqual(ws.max_column, 7)
            self.assertEqual((ws["A6"].value, ws["B6"].value, ws["C6"].value, ws["F6"].value, ws["G6"].value),
                             (1, "33KV", "Campo 04", "A", "204 A"))


if __name__ == "__main__":
    unittest.main()
